from pathlib import Path
import json, html, re, base64, shutil, zipfile, hashlib
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT=Path(__file__).parent
OUT=ROOT/'开发定稿'
OUT.mkdir(exist_ok=True)
ASSETS=OUT/'设计资源'
ASSETS.mkdir(exist_ok=True)
UI=Path('/Users/zhengjunfeng/Desktop/temp/player_ui')
blocks=[]
def h(x):blocks.append(('h',x))
def s(x):blocks.append(('s',x))
def p(x):blocks.append(('p',x))
def b(x):blocks.append(('b',x))
def t(head,rows):blocks.append(('t',(head,rows)))
def c(x):blocks.append(('c',x.strip()))
def img(name,caption):
    shutil.copy2(UI/name,ASSETS/name)
    blocks.append(('image',(name,caption)))

tasks=json.loads((ROOT/'交付文件/Mac播放器_版本TODO.json').read_text())
sources=json.loads((ROOT/'交付文件/调研来源.json').read_text())
sources.extend([
dict(id='S33',title='Apple TV Like Player 界面参考与 MIT 许可',url='https://github.com/doraFX/apple-tv-like-player',type='已确认界面参考'),
dict(id='S34',title='mpv OpenGL 渲染接口',url='https://github.com/mpv-player/mpv/blob/master/include/mpv/render_gl.h',type='上游接口源码'),
dict(id='S35',title='mpv 实际许可与构建条件',url='https://github.com/mpv-player/mpv/blob/master/Copyright',type='上游许可说明'),
dict(id='S36',title='GRDB SQLite Swift 工具库',url='https://github.com/groue/GRDB.swift',type='上游项目'),
dict(id='S37',title='Apple 网站分发公证流程',url='https://developer.apple.com/documentation/security/notarizing-macos-software-before-distribution',type='苹果开发文档'),
dict(id='S38',title='NSViewRepresentable',url='https://developer.apple.com/documentation/swiftui/nsviewrepresentable',type='苹果开发文档'),
])
source_map={x['id']:x for x in sources}
changes={
'A01':('冻结产品范围与五条任务流程','使用本定稿边界；设计风格已获用户认可；形成五条可重复的用户验收任务'),
'A02':('将已确认 UI 转成原生组件规格','以交付设计稿为视觉基准；核对三种 Mac 窗口、状态、键盘与辅助功能；不用重新设计方向'),
'A05':('建立 Swift 工程与固定依赖','macOS 14+ arm64；锁定 Xcode、Swift、libmpv、FFmpeg、GRDB 版本和校验值；可干净构建'),
'B03':('实现已确认的播放控制界面','中央播放与前后 10 秒，底部时间线音量全屏，顶部字幕音轨倍速队列；静止自动隐藏规则通过'),
'B07':('实现数据层与播放进度','SQLite 迁移与事务；按文件保存进度、轨道与音量；退出续播误差目标小于 2 秒；记录默认每 5 秒保存'),
'B10':('实现首版设置和简中英文文案','浅色默认；播放、字幕、数据与更新设置完成；键盘与 VoiceOver 可访问；系统黑暗模式的替代不假称已适配'),
'C01':('实现来源与文件身份模型','同一来源文件不因简单更名丢进度；内容替换不继承旧记录；保存卷标识和相对路径'),
'C05':('实现已确认的首页和侧栏','首页继续观看与最近加入；视频库、文件夹、收藏、搜索、设置齐备；原生三尺寸视觉验收'),
'C06':('实现真实文件夹树和浏览密度','不是只有来源根目录列表；任意子目录可进入和返回；网格列表切换、面包屑和滚动恢复可用'),
'C07':('实现搜索收藏和观看标记','1 万条记录可检索；收藏与单项已看可撤销；批量已看后置；原视频不改变；支持手工类型标记与筛选，未知内容保留原文件名'),
'C08':('完成播放与浏览无重载切换','默认返回浏览继续播放并显示底部当前播放栏；原滚动位置、选中项和同一内核会话保留'),
'C10':('完成键盘辅助使用和大库性能','三种 Mac 窗口、VoiceOver、减少动态效果与透明度通过；1 万条索引的性能有测量记录'),
'D01':('实现作品分类与批量已看','电影剧集个人视频分类可手工修正；批量已看有撤销；不以早期示例旅行标签作为自动分类规则'),
'E05':('实现有限外观参数及深色主题','默认保留已确认浅色；深色经过完整对比度和状态验收才启用；密度、点缀色、控件可见性可调'),
'G05':('完成辅助功能与三尺寸 UI 回归','960×640、1280×800、1440×900；键盘、VoiceOver、浅色深色、减少动态与透明度通过'),
}
version_names={'A':'v0.0','B':'v0.1','C':'v0.2','D':'v0.3','E':'v0.4','F':'v0.5','G':'v0.9','H':'v1.0'}
for task in tasks:
    if task['id'] in changes:task['title'],task['acceptance']=changes[task['id']]
    task['version']=version_names[task['id'][0]]
    task['dependsOn']=list(filter(None,task.pop('depends').split(',')))
    task['status']='未开始'
    task['evidenceRequired']='实现记录、相应测试结果和人工验收；勾选规划不替代完成证据'
    task['designBaseline']='UI设计稿.html' if task['id'] in ['A02','B03','B05','B09','B10','C05','C06','C08','G05'] else None
for task in tasks:
    if task['id']=='A03':task['priority']='P1';task['acceptance']='优先复核 IINA 和 Owl 的五任务路径；其他竞品按可安装条件补充，不阻塞播放 POC'
    if task['id']=='A10':task['dependsOn']=[x for x in task['dependsOn'] if x!='A03']
    if task['id']=='B08':task['acceptance']='原生全屏、Esc、空格、方向键通过；Now Playing 和媒体键在取得命令归属时响应，不注册全局监听抢占其他应用'
extra=[
('C15','补齐缩略图与文件信息的真实数据','时长、分辨率、编码和音轨来自探测；无数据用未知；不得硬编码 4K、H.265、AAC；收藏封面为真实抽帧',['C04','B05']),
('C16','补齐当前播放栏与恢复播放流程','浏览时底部栏显示真实会话；关闭/切换页面不重复发声；从头播放和继续播放明确',['C08','B07']),
('C17','完成子文件夹和多文件来源视图','根来源与真实子目录区分；离线、扫描中、空目录、无搜索结果状态完整',['C03','C06']),
('C18','完成媒体库导出与恢复','导出含 schemaVersion、来源、逻辑 ID、进度、收藏、设置；重新导入匹配现有文件且可回滚；不含视频',['C09']),
]
for tid,title,accept,deps in extra:tasks.append(dict(id=tid,title=title,acceptance=accept,version='v0.2',dependsOn=deps,size='M',priority='P0',status='未开始',evidenceRequired='真实数据与操作路径验收',designBaseline='UI设计稿.html'))
for task in tasks:
    if task['id']=='C14':task['dependsOn']=list(dict.fromkeys(task['dependsOn']+['C13','C15','C16','C17','C18']))
tasks.sort(key=lambda x:(list(version_names.values()).index(x['version']),int(x['id'][1:])))

h('1 开发决策与产品边界')
p('开发一款仅面向 macOS 的本地视频播放器。采用已经确认的浅灰蓝、白色、深海军蓝与圆润控件风格；播放页参考 Apple TV Like Player 的画面布局和半透明控制层。首个公开版为 v0.2，先完成可靠播放、目录浏览、收藏、搜索和续播，再增加内容整理、播放方案及标准网络流。')
p('工程暂定名 Mori。产品定位是“可以随手打开，也可以认真收藏的 Mac 本地播放器”。UI 方向已确认；真实播放内核、原生界面、安装包和媒体库尚待开发。公开发行前确认展示名称、图标和实际 GitHub 所属账户对应的 Bundle ID，不以此阻塞内部工程。')
t(['事项','开发定稿'],[
['平台','macOS 14 及以上；先发布 Apple Silicon arm64；Intel 在具备真机测试后另立版本'],
['应用形态','Swift 原生桌面应用；不以浏览器或 WebView 作为主 UI 和广泛格式播放底座'],
['语言和内核','Swift 6；SwiftUI 加 AppKit；libmpv 加固定 FFmpeg 依赖；首版不引入 Rust'],
['服务端','无需自建业务服务器、账号或数据库服务；核心播放、索引、收藏、进度离线可用'],
['分发','静态官网与 GitHub Releases 下载；Developer ID 签名、公证、自带依赖的 DMG；Sparkle 2 更新'],
['媒体策略','只读引用本地原文件；播放不转码；加入媒体库不复制视频；不提供原视频删除移动重命名'],
['开源','推荐 GPL-3.0-or-later；发行时按实际二进制依赖核验许可证和对应源码'],
['首版视觉','默认浅色；圆润但紧凑；播放器画面优先；高级操作按需展开']])
p('无需业务服务器不等于应用永不联网。更新检查读取静态 appcast 和发布包；v0.3 可选第三方元数据，v0.5 标准网络视频都由桌面应用直接访问。它们默认不影响本地播放，关闭后仍能使用本地核心功能。')
p('不纳入项目：此前的监控加密文件、解密工具、跨平台、商业会员、App Store 上架、云盘聚合、内容下载、视频编辑、AI 清晰度增强和服务器转码。标准 H.265 视频按普通编码测试和支持。')
s('差异化与验证')
t(['面对的选择','本项目要完成的具体体验','验证方式'],[
['IINA 与 VLC','保留 Finder 直接打开的效率，同时把播放器、目录、收藏和返回路径做成统一应用','同一文件打开、返回目录、下次续播的任务测试'],
['Infuse 与 VidHub','优先一台 Mac 的原文件夹和个人视频；不要求先联网匹配为电影才能好用','断网导入、个人视频检索、离线硬盘恢复'],
['Owl 与其他新原生播放器','在目录浏览基础上打磨稳定的播放与返回、可调信息密度以及后续可撤销播放方案','同任务完成率、错误次数、两周持续使用']])
p('上述是产品实现目标和迁移假设，不能宣称竞品均不支持这些能力。漂亮、开源和格式覆盖已有成熟选择，差异应落实到完整操作路径。市场依据保留在随包的研究资料中；产品调研日期为 2026 年 10 月 12 日，接口与交付基线补充核对日期为 2026 年 10 月 13 日。[S03][S04][S06][S08][S11]')

h('2 首版 PRD 与页面验收')
s('首个公开版 v0.2')
p('目标用户是从 Finder、下载目录或外接硬盘观看视频的人，以及希望保留原目录又能收藏和续播的用户。首次打开提供“打开视频”和“添加文件夹”；没有账号、网络或导入全库的前置要求。')
t(['需求编号','能力','必须满足的行为'],[
['R01','打开与队列','Finder 双击、打开面板、拖放文件和文件夹；中文与空格路径；多文件自然排序；当前会话可替换或加入队列'],
['R02','播放控制','播放暂停、前后 10 秒、拖动定位、音量静音、倍速、原生全屏；状态来自真实内核'],
['R03','字幕音轨','内嵌和外部字幕、语言、延迟、字号、音轨切换；无轨道时如实显示；单字幕为首版范围'],
['R04','观看进度','每 5 秒及暂停、切换、退出时保存；从头/继续明确；不把手动停止当自然播完'],
['R05','来源和子目录','添加来源后保留全部目录层级；子目录可进入返回；枚举与缩略图分阶段；可取消扫描'],
['R06','媒体浏览','首页继续观看和最近加入；视频库网格与列表；舒适与紧凑密度；原文件名可辨认'],
['R07','收藏与搜索','收藏、取消及撤销；标题、原文件名和文件夹名检索；首版手工类型标签，不自动识别旅行内容'],
['R08','浏览与播放共存','返回浏览默认继续播放，显示底部当前播放栏；同一播放会话，不重新解码；返回位置保持'],
['R09','离线与出错','拔盘只标离线；不删除条目；重连核对恢复；损坏、缺失、权限和解码错误分别给恢复动作'],
['R10','设置与本地数据','播放字幕常用设置、缓存清理、记录清除、来源级不记历史、导出恢复；不改原视频'],
['R11','发布与可用性','简中英文、键盘和 VoiceOver；干净 Mac 可安装；首次运行无需 Homebrew 或额外解码包']])
s('页面与完整状态')
t(['页面','主内容和动作','状态与版本'],[
['播放页','返回、中央播放与前后 10 秒、时间线、音量、全屏、字幕音轨、倍速、队列、收藏、文件信息','加载、播放、暂停、定位、结束、失败；v0.1'],
['首页','继续观看、最近加入、添加文件夹、底部当前播放栏','首次空态、已有内容、扫描中；v0.2'],
['视频库','全部视频、网格列表、排序、密度、详情、手工筛选','空库、筛选空、探测中、缺图；v0.2'],
['文件夹','来源列表、真实子目录、面包屑、数量、进入和返回','在线、离线、无权限、失效路径、扫描取消；v0.2'],
['收藏与搜索','收藏视频、取消、关键词、结果数量','没有收藏、没有结果、搜索中；v0.2'],
['设置','外观、播放、字幕、存储、关于与更新','已保存、失败、危险操作确认；v0.1 至 v0.2'],
['添加与文件信息','选择目录、包含子目录、扫描进度、取消；真实文件属性','设计稿的示例目录与属性须换成系统面板和探测结果；v0.2'],
['后续详情','电影、剧集、多版本、可选在线信息','v0.3；不阻塞本地个人视频']])
s('关键交互规则')
b('浏览卡片：单击选中；双击或明确的播放按钮进入播放。信息按钮打开文件信息。原型中的整卡点击播放可用于演示，正式 Mac 行为以上述规则为准。')
b('播放浮层：鼠标静止约 2.5 秒后隐藏；暂停时保留；菜单展开、键盘焦点、拖动、音量操作时不隐藏。中央控件只在浮层可见时出现，避免始终遮挡画面。')
b('返回浏览：保持播放会话和声音，显示当前播放栏，恢复来源、子目录、搜索词、筛选、排序、滚动与选中项。设置可选择返回时暂停，默认继续。')
b('Esc：先关闭弹层，再退出系统全屏，再返回浏览。空格播放暂停，左右方向键前后 10 秒；文本框和可编辑控件取得焦点时不拦截文本操作。')
b('拖动进度：拖动时先本地显示目标时间，合并中间定位意图，释放后提交最终精确定位；原本暂停时定位后继续暂停；旧请求不得覆盖新位置。')
b('完成判断：普通长视频建议达到 95% 或剩余不超过 120 秒时标已看；不足 10 分钟的短片建议只用 95%；短于 30 秒不自动续播。用户手工标记优先；阈值单测并可后续调整。')
b('从媒体库移除只解除索引和引用；清理缓存只删除应用自己的缓存；清除记录需确认。公开版不提供删除原视频入口。')

h('3 已确认 UI 与原生设计规范')
p('视觉采用柔和极简风格。浅灰蓝侧栏和白色内容区承载深海军蓝文字与主按钮，暖杏色只用于少量进度和选中提示。播放页以视频画面为主体，半透明圆形控制组、细时间线和少量图标沿用用户给出的播放器参考。[S33]')
img('player-final.png','已确认播放页基准  示例画面和视频属性仅用于展示布局')
img('home-desktop.png','已确认首页基准  色彩来自视频内容  不使用网站式营销大封面')
p('随包 UI设计稿.html 可切换全部已有页面和异常状态，用于确认布局和交互意图。它不是可播放广泛格式的应用，也不是生产 WebView 模板；素材、视频名、属性、扫描和轨道均为示例。设计稿中的“查看界面”切换器、窗口红绿灯装饰和演示文案不进入产品。正式窗口使用 macOS 原生标题栏和交通灯。')
s('颜色 排版与间距')
t(['类别','定稿初值','实施规则'],[
['颜色','文字 #141D2D；次要文字 #647084；内容 #FFFFFF；侧栏 #F0F4F8；分隔 #E5EAF0；点缀 #EFB98F','以设计规范.json 为唯一 token 来源；暖色不承载小字号正文'],
['字体','系统 SF 与 PingFang；页标题 27 pt；节标题 18 pt；正文 14 pt；元信息 12 pt','使用可缩放字体样式；时间数字等宽；图标优先 SF Symbols'],
['窗口与布局','最小 900×600 pt；基准 960×640、1280×800、1440×900；侧栏 176 pt；内容边距 30 pt','宽度不足先收侧栏；播放画面保留原始比例，默认留黑边而非裁切'],
['圆角','窗口视觉 20 pt；内容缩略图 14 pt；弹层 16 至 22 pt；主按钮胶囊','系统标题栏与实际窗口外轮廓跟随平台；视频全屏不加装饰边框'],
['网格','继续观看卡片 16:9；个人视频 16:10；后续电影海报 2:3；卡片间距 16 至 24 pt','只裁切缩略图；原视频播放不继承封面裁切；密度变化保持滚动锚点'],
['播放控制','中央播放 88 pt；两侧 62 pt；间隔 26 pt；图标约 39 和 27 pt；底部横向边距 22 pt','按 Mac 窗口自适应，控制命中区至少 32 pt，常用目标建议 44 pt'],
['动效与材质','交互 160 至 240 ms；玻璃仅用于视频浮层；减少透明度时用高对比实色','尊重减少动态效果；不做持续装饰动画；聚焦与选中不能只靠颜色']])
p('以上为自己的组件参数，不是 Apple 内部设计参数。网页设计稿的 CSS px 可作为逻辑尺寸起点，原生以 pt 实现并在 Retina 真机核对。默认浅色已确认；深色完整方案放在 v0.4，完成对比度和状态验收前不提供会产生黑底白字冲突的“跟随系统”。')
s('原型未覆盖但首版必须补齐')
b('真实子目录页面和面包屑，而不是只显示来源根文件夹；多层目录返回后恢复选中项。')
b('浏览时底部当前播放栏；设计稿切页面暂停是演示简化，正式应用按 R08 实现。')
b('真实加载缓冲状态、音频设备切换、无字幕和无音轨、未知视频属性；不能保留示例 4K 和 H.265 标签。')
b('真正的系统打开面板、全屏、文件权限与扫描；“模拟重新连接”等演示按钮全部移除。')
b('真正的空库、空文件夹、搜索无结果、失效来源和迁移失败恢复页面；普通文件拖入即播放不必入库。')

h('4 技术架构与播放实现')
s('技术栈与工程结构')
t(['模块','选型','约束'],[
['应用','Swift 6 与当前稳定 Xcode；macOS 14+ arm64','开启严格并发检查；不要把 C 指针用 unchecked Sendable 随意传递'],
['UI 与窗口','SwiftUI 组织页面；AppKit 管理 NSWindow、全屏和视频宿主','NSViewRepresentable 包装稳定宿主；媒体会话不随 View 重算销毁 [S38]'],
['播放','libmpv，使用固定版本头文件和一致二进制','MPVAdapter 隔离 C API；不加载用户全局 mpv.conf 或第三方脚本'],
['渲染','先验证 libmpv render API 的 OpenGL 嵌入','公开 API 提供 OpenGL 与软件路径；不预设可直接嵌入 Metal [S19][S34]'],
['数据','SQLite 加 GRDB；显式 repository 与 migration','SPM 固定可构建 release；WAL、短事务；UI 不直接执行 SQL [S36]'],
['探测与缩略图','随 App 打包 ffprobe 和 ffmpeg 辅助进程','固定程序路径和参数数组；不依赖用户已有 FFmpeg；取消与超时必须回收进程'],
['更新','Sparkle 2 与静态 HTTPS appcast','实际固定稳定版本；更新 EdDSA 签名与 Apple 签名分别验证 [S27]']])
c('''MoriApp                 应用生命周期和依赖注入
AppShell                导航状态 WindowCoordinator 当前播放栏
DesignSystem            颜色 字体 间距 控件 token
PlayerFeature           PlaybackViewModel 播放浮层 队列 轨道菜单
LibraryFeature          首页 视频库 文件夹 收藏 搜索
SettingsFeature         播放 字幕 存储 更新
Domain                  ID 值类型 状态 能力 命令
Playback                PlaybackCoordinator MPVAdapter EventPump
Rendering               RenderSurface AppKitVideoHost RenderExecutor
Persistence             GRDB Repository Migration Backup
FileAccess              SourceAccess VolumeMonitor ScanService
MediaAnalysis           ProbeService ThumbnailService Cache
Distribution            UpdateService Diagnostics
Tests                   Unit Integration UI MediaFixtures
Build                   依赖锁 构建脚本 签名 公证 许可清单

UI → ViewModel → Coordinator / Repository
Coordinator → PlaybackEngine → MPVAdapter → libmpv
RenderSurface → AppKit 宿主 → mpv_render_context
ScanService → Probe / Thumbnail → SQLite / Caches
''')
p('首版采用一个应用工程和清晰的 Swift 模块目录；确有隔离需求再抽成 local Swift Package。避免为个人项目拆出微服务、通用插件内核或多个重复播放器。数据库、播放和渲染的职责要隔开，UI 不持有 mpv_handle 或数据库连接。')
s('播放契约与事件')
c('''PlaybackEngine  仅描述业务命令与事件
  open(OpenRequest)              异步加载
  setPaused(Bool)                显式暂停意图
  seek(seconds, mode, seekID)     合并定位
  setVolume(Double) setMuted(Bool)
  setRate(Double)
  selectAudio(TrackID)
  selectSubtitle(TrackID?) addSubtitle(URL)
  stop() shutdown()
  events() → AsyncStream<PlaybackEvent>

OpenRequest: requestID source resumeAt selectedTracks
EventEnvelope: sessionID requestID sequence payload
PlaybackEvent: loading ready position tracksChanged
               pauseChanged buffering seekCompleted ended failed
PlaybackCapabilities: seekable live subtitles hardwareDecode hdrMode
RenderSurface: 平台对象 单独管理 不跨 actor 发送
''')
p('这是一份接口契约，不是可编译实现。用不可变 Sendable 值传递事件；UI 的状态模型位于 MainActor。每次 open 创建新 requestID，序列号保持有序；丢弃旧 request 的晚到结果；mpv native track ID 只在当前文件和会话内解释，轨道选择持久化需保存语言、类型和标题等匹配提示。')
c('''用户意图                 paused / playing
媒体阶段                 idle → opening → ready → ended / failed
暂态                     buffering / seeking

open 新文件              保存旧进度 取消旧加载 更新 requestID
FILE_LOADED              读取轨道与真实能力 应用续播和用户意图
seek 完成                只接收当前 seekID 保持原暂停状态
END_FILE                 区分自然结束 停止 错误 替换
自然结束                 根据自动下一项与队列决定后续
返回浏览或系统全屏        只改变布局 不重建内核
退出                     停止投递 解除回调 释放渲染 销毁内核
''')
s('执行域与渲染约束')
p('内核控制命令和事件消费进入专用串行执行域；C wakeup 回调仅通知事件泵。渲染使用独立执行域，确保一次只调用一个相关 mpv_render API，且 OpenGL context 在调用线程 current。Swift actor 保证隔离，不保证固定 OS 线程或 OpenGL context，不能用一个 actor 名字代替上下文管理。[S19][S20]')
p('渲染线程不能同步等待不安全的内核命令，也不能和主线程形成锁等待环。正确实现的顺序是创建宿主与 context、初始化 render context、再加载媒体；释放时反向处理并解除回调。SwiftUI 只接收低频进度与菜单状态，目标约 10 Hz；视频帧不转为每帧 SwiftUI Image，不在主线程拷贝全尺寸 RGBA。')
p('首个 POC 用 OpenGL 路径验证 SDR 播放、半透明覆盖、缩放、Retina drawable 尺寸、最小化、锁屏和系统全屏。先避免窗口间重新挂载复杂度；主窗口内切换浏览时保留同一会话与宿主。小窗和系统 PiP 分开后置，绝不把新的 mpv 实例当作窗口迁移的解决办法。')
p('OpenGL 是当前可验证的嵌入路线，不是长期承诺。若两周 POC 仍无法达到稳定叠层、颜色或窗口生命周期要求，先阻止媒体库扩展，再单独评估 libVLC 或 AVFoundation 等路线；不要边开发 v0.2 边默默重写内核。')
s('画质 声音和格式支持')
p('原视频只读，播放过程解码和显示，不执行重新编码；关闭默认锐化、降噪、插帧和裁切，不修改宽高比。清晰度仍取决于原片、显示缩放、色彩范围和显示器。缩略图压缩不影响播放视频；截图也不能证明 HDR 输出正确。')
p('优先验证 H.264、H.265 Main 与 Main10 的 SDR 路径；硬件解码通过实际 mpv 能力和机器表现决定，失败可降级软件解码并展示性能提示。硬解不代表所有 H.265 profile 都硬解，也不代表画质必然优于软件解码。')
p('首版只承诺经测试的 SDR，以及通过验证时可用的 HDR 到 SDR 映射；真实 HDR/EDR 输出另立 POC。未通过映射就把对应样本列为限制，不显示“原生 HDR”标签。AAC、PCM、FLAC 等解码为系统音频输出；多声道下混、蓝牙和输出设备变化单测，AC3/DTS 解码不等于 Atmos 或 TrueHD 位流直通。[S21][S32]')
t(['P0 媒体集合','样本和验收'],[
['封装','MP4、MOV、MKV、WebM、AVI、TS 的常用组合；缺索引、损坏文件单列'],
['视频','H.264；H.265 Main/Main10；VP9；AV1 常用样本；硬解与软解分别标识'],
['音频','AAC、MP3、PCM、FLAC、Opus；AC3/EAC3 和 DTS 解码样本；切轨不丢声'],
['字幕','SRT、ASS、VTT；MKV 内嵌；中文编码；复杂 ASS；PGS 等图片字幕有样本才承诺'],
['画面与时序','旋转 MOV；23.976/24/25/30/60 fps；可变帧率；4K60、外接显示器分别测'],
['硬件操作','睡眠恢复、蓝牙切换、拔盘、100 次切换文件与窗口、2 小时播放']])
p('能力清单由随包二进制实际构建选项和测试结果共同产生。对外描述“支持常见格式，详见兼容清单”，不承诺“所有格式”。ISO/BDMV 菜单、DRM、特殊 12-bit/4:2:2 和所有网络站点不属于 v1 保证。')

h('5 本地数据与文件安全')
s('实体与演进')
t(['实体','v0.2 字段与约束','后续演进'],[
['Source','UUID、rootURL/bookmark、volumeID、state、scanGeneration；状态 online/offline/permissionDenied','目录级可选元数据；网络源单独类型'],
['MediaItem','UUID、displayTitle、manualType、manualOverrides；允许 unknown/personal','v0.3 引入 movie/episode 分类、作品匹配'],
['MediaFile','UUID、itemID、sourceID 可空、relativePath、resourceID、size、mtime、fingerprint、availability；source 内路径唯一','item 一对多 file；临时打开也可保存本地续播身份'],
['Track','fileID、kind、language、codec、title；真实轨道数据','nativeTrackID 不用作跨会话全局 ID'],
['PlaybackState','fileID 唯一、position、duration、watchedOverride、trackHints、updatedAt','进度按具体文件，不盲目套用到不同版本'],
['Favorite','itemID 唯一、createdAt','v0.3 需要集合时再加 Collection 和 Member'],
['CacheEntry','key、kind、bytes、lastAccess、fingerprint；实际图片在 Caches','LRU 与失效版本；缓存可重建'],
['后置实体','不在首版预建账号和服务端表','MetadataMatch v0.3；Profile、Bookmark v0.4']])
p('默认同一内容文件只有一个稳定 ID。绝对路径、inode、大小修改时间和抽样指纹都不能单独证明同一文件；匹配按卷和 resource 标识、相对路径、属性、后台抽样指纹分层组合。出现冲突时保留两个条目让用户确认，不自动合并而覆盖进度。内容替换后重新探测并生成新身份。')
p('SQLite 使用事务和迁移版本，Source/MediaFile 通过外键关联；必要查询建立 source/path、item、updatedAt 索引。WAL 允许读取与短写入并行；扫描分批提交。中文搜索先实现受限 LIKE 和标准化文件名查询，绑定参数并转义通配符；不能把默认 FTS 分词器当作中文子串搜索的完整解决方案。1 万条压测不达标再引入匹配索引。')
s('渐进扫描和目录重连')
b('第一阶段枚举文件名和基础属性并立即显示；第二阶段获取真实媒体属性；第三阶段为可见卡片和近期内容优先抽帧。媒体库空态与探测中状态不同。')
b('目录事件只作为需要核对的提示，去抖后增量对账；挂载变化与事件丢失执行完整核对。只有来源在线、权限有效且整轮枚举成功，才允许把本轮未见项标 missing。扫描取消和拔盘不能批量删索引。')
b('避免循环符号链接；嵌套来源明确去重；文件写入中先等待稳定属性再探测。缩略图最多两个后台进程，前台播放时减为一个或暂停；每个任务有取消和超时，输出限制和缓存上限初值 1 GB。')
b('非 App Store 首版不启用 App Sandbox，仍尊重 macOS 隐私与文件权限。普通 URL bookmark 可用于位置恢复；不要把 security-scoped bookmark 当成非沙盒所有路径必须步骤。未来启用沙盒时独立补访问授权生命周期。')
b('重定位先验证用户所选目录和原来源的内容标识；确认后事务更新 root，保留逻辑 ID、收藏和进度。错误路径与不完整扫描不能覆盖有效来源。')
s('存储位置与恢复')
c('''Application Support/<bundle-id>/library.sqlite
Application Support/<bundle-id>/settings.json
Application Support/<bundle-id>/Backups/
Caches/<bundle-id>/thumbnails/
Caches/<bundle-id>/probe/
Logs/<bundle-id>/                 有界本地诊断日志
Keychain                          仅后续网络凭据
''')
p('数据库迁移前创建可恢复备份；SQLite WAL 模式用数据库备份 API 或受控 checkpoint/停写快照，不在有并发写入时只复制 .sqlite 主文件。迁移失败保留原库、进入恢复流程，不静默清空重建。导出格式含 schemaVersion、逻辑 ID、来源位置提示、收藏、进度和配置，不含原视频和可执行脚本。')
p('默认不发送遥测。来源可选择不记观看历史；诊断导出先展示字段并剔除文件全路径、视频名与后续 URL token。应用清理和卸载说明区分本地记录、可重建缓存与原始媒体。')

h('6 分版本路线与开发任务')
p('全部工程任务以开发任务.json 为准，包含固定 ID、版本、优先级、依赖、规模和逐项验收。UI 风格确认是设计决策，不代表原生实现或兼容测试完成。任务当前均标为未开始，需提交实现、测试和验收证据后再改变状态。S 约半天至 1 天，M 约 1 至 2 天，L 约 3 至 5 天，作为相对规模。')
t(['版本','用户结果','发布门槛','规划时间'],[
['v0.0','最小原生 H.265 播放与已确认 UI 叠层','30 样本、正确 SDR、渲染生命周期、依赖可复现；失败则停扩展','约 2 周'],
['v0.1','可用的本地播放器','打开播放、轨道字幕、续播、全屏、错误恢复；60 样本与内部包','约 3 周'],
['v0.2','第一份公开体验版','首页、视频库、目录、收藏、搜索、无重载返回、签名公证更新；用户任务验收','约 3 至 4 周'],
['v0.3','作品与电影剧集整理','本地封面和 NFO、手工优先、可选在线信息；断网仍可用','约 2 至 3 周'],
['v0.4','可玩性和完整外观','快捷键、可撤销方案、双字幕、书签循环、时间轴预览；深色主题','约 2 至 3 周'],
['v0.5','标准网络视频','HTTP/HLS/DASH 样本、能力 UI、认证与取消；不保证动态 ABR','约 2 至 3 周'],
['v0.9','发布候选版','120 样本、历史迁移、硬件与干净安装回归；冻结功能','约 2 周'],
['v1.0','稳定版本与维护基础','阻断与严重缺陷清零；安装包、更新、源码和许可对应','约 1 周']])
p('v0.2 按约 8 至 10 个全职工作周规划；完整 v1.0 约 22 至 28 周，含约 30% 兼容和返工缓冲。新增四项补齐任务以并入相关任务为主；如果真实子目录、导出恢复超出原规模，需要重新估算，不能用版本号倒逼删掉数据安全验收。上述不是发布日期承诺。')
s('任务与需求对应')
t(['验收域','任务 ID','关键结果'],[
['播放与格式','A04 至 A10、B01、B04、B11','样本、状态机、渲染、定位和颜色通过'],
['用户控制','B02、B03、B05、B06、B08 至 B10','文件打开、浮层、字幕、队列、全屏与错误恢复'],
['进度和浏览','B07、C01、C05 至 C08、C16','身份与续播、首页、收藏搜索、播放栏与返回'],
['目录和原始媒体','C02 至 C04、C15、C17','真实子目录、扫描取消、离线恢复、真实属性'],
['数据恢复和隐私','C09、C18、G03','导出恢复、迁移备份、隐私来源'],
['发布','B12、C11 至 C14、G06 至 G08、H01 至 H05','可下载安装、升级、体验测试和稳定发布']])
s('第一周和第二周')
t(['时间','具体产出','完成定义'],[
['第 1 天','建原生工程、模块目录、依赖锁和许可目录；导入设计 token','可以干净构建；保留 UI 参考，不继续画新的大封面'],
['第 2 天','libmpv 最小打开与事件桥；30 个授权基础样本 manifest','H.264 与 H.265 的声音画面都可验证'],
['第 3 天','AppKit 稳定视频宿主和 SwiftUI 控制叠层','播放暂停、前后跳转、真实时间与音量有响应'],
['第 4 天','窗口缩放、Retina、全屏、100 次切换、停止释放','无重复声音、渲染崩溃或主线程阻塞'],
['第 5 天','SDR、硬软解路径、首帧和同步基线；记录 POC 风险','把能通过与不能通过的能力写清；决定第二周重点'],
['第二周','补齐失败恢复、生命周期、打包依赖和 POC 证据','A10 通过才进入 v0.1；先别做电影元数据和云端连接']])
p('开发顺序先完成真实播放纵切片，再接数据和目录，最后接完整媒体库。UI 可用假数据并行搭建组件，但合入主线前必须替换示例属性和模拟流程。单个开发任务至少包含需求 ID、成功与失败路径、日志和验收样本。')

h('7 质量指标与发布流程')
s('测试目标')
t(['指标','v0.2 验收目标','测量条件'],[
['冷启动','P95 不超过 2 秒','首轮系统验证已完成；本地 SSD；重复 20 次'],
['首帧','1080p H.264/H.265 P95 不超过 1.5 秒；4K 不超过 2.5 秒','记录机器、编解码和渲染路径；首声单独记录'],
['定位','常规本地样本 P95 不超过 800 ms','至少 20 点；缺索引和长 GOP 单列'],
['同步','30 分钟无持续漂移；目标误差不超过 80 ms','专用同步样片与明确测量；听感不替代证据'],
['媒体库','1 万条索引检索 P95 不超过 200 ms；热缓存滚动 95% 帧不超过 16.7 ms','中文搜索、扫描同时运行分别测'],
['内存与能耗','1 万条浏览目标 350 MB；1080p 播放目标 600 MB；对照同设置 IINA 的能耗基线','4K/HDR 单独报告；100 次切换无持续增长'],
['稳定性','所有 P0 样本通过；100 次切换；2 小时播放；拔盘和睡眠回归','出现严重故障先修复，不扩展功能']])
p('指标是待实测的验收目标，不是当前性能结果。v0.0 建 30 样本，v0.1 扩到 60，候选版至少 120；manifest 记录容器、codec、轨道、时长、许可来源、校验值和预期结果。不要把商业电影直接提交公开仓库。')
p('单元测试覆盖身份匹配、自然排序、已看判断、SQL 迁移、进度与方案校验；集成测试使用真实 libmpv 事件；UI 测试覆盖五条主路径；颜色、多屏、蓝牙、声音和能耗使用真机。原型按钮能点击不等于播放器验收通过。')
s('网站发布和自动更新')
c('''依赖锁与源码归档 → 测试 → Release 构建
 → 收集 dylib 和 helper 设置应用内 rpath
 → 核验架构 依赖 构建配置 许可证
 → 从内到外签名库 helper 和 App
 → Hardened Runtime 与最少必要 entitlement
 → 打包 DMG 与更新归档
 → Apple 公证并 staple 票据
 → Gatekeeper 验证 SHA256 Sparkle EdDSA 签名
 → 干净 Mac 从浏览器下载安装和升级
 → 发布具体版本与 appcast
''')
p('官网静态托管即可，下载指向具体 Release，明确最低系统、arm64、校验值、已知限制、源码和反馈入口。Apple 公证与 Developer ID 用于正常直接分发；不以让用户关闭 Gatekeeper 或执行删除隔离属性命令作为安装流程。[S25][S37]')
p('stable/beta 分两个更新渠道；CFBundleVersion 单调增加；播放时只提示更新，不强制退出。签名密钥放专用 Keychain 或受保护的 CI secret，不进仓库和日志。安装包自带全部播放依赖，用户无需本机 FFmpeg。Sparkle 的更新签名不能替代 Apple 代码签名。[S27]')
s('许可证与参考仓库')
p('应用推荐 GPL-3.0-or-later。mpv 默认 GPLv2-or-later，关闭 GPL 文件的构建可为 LGPL，但链接依赖仍影响最终许可；FFmpeg 的构建配置也需逐项核验。首版不使用 nonfree 组件。每个发布版本提供对应源码、补丁、构建脚本、配置、依赖版本和第三方 notices。[S35]')
p('doraFX/apple-tv-like-player 是 MIT 许可的 HTML5 控件参考，不解决原生解码或系统格式支持。本定稿采用其交互方向，用 SwiftUI/AppKit 重新实现；如之后复制其 CSS/JS 或资源，保留对应许可证和版权声明。开源代码许可不自动授予 Apple 品牌、他人海报或商业视频素材的使用权。[S33]')

h('8 开发启动与完成标准')
p('第一步不是继续做界面展示，而是在原生窗口中可靠播放一个真实 H.265 文件，验证声音、颜色、定位、停止释放和控制叠层。这个纵切片完成后，再按 B 组和 C 组任务接入数据与媒体库。')
t(['开发阶段','完成标准'],[
['一个任务','代码与对应需求一致；正常、失败、取消路径可复现；相关测试通过；没有硬编码演示属性'],
['一个版本','P0 任务和样本全部通过；数据迁移与原文件不变；已知限制公开；真实 UI 和安装路径可用'],
['首个公开版','从网站下载后可安装；Finder 打开视频；回到文件夹；切字幕；重启续播；拔盘重连不丢记录'],
['v1.0','兼容矩阵、签名更新、源码许可、用户说明与实际安装包一一对应；阻断和严重问题清零']])
p('本开发定稿版本 1.1，日期 2026 年 10 月 13 日。产品范围、浅色视觉方向、首版功能和实施路线已确定。早期市场调研作为背景资料保留；涉及 UI、版本范围、首版深色主题和交互规则时以本定稿为准。')

tokens={
'schemaVersion':1,'product':'Mori','status':'用户已确认视觉方向，原生实现待开发','platform':'macOS 14+ arm64','units':'pt for native layout',
'colors':{'textPrimary':'#141D2D','textSecondary':'#647084','surface':'#FFFFFF','sidebar':'#F0F4F8','separator':'#E5EAF0','warmAccent':'#EFB98F','playerFallback':'#111722'},
'typography':{'family':'System SF / PingFang','pageTitle':27,'sectionTitle':18,'body':14,'metadata':12,'timeDigits':'monospaced'},
'spacing':{'sidebarWidth':176,'contentInset':30,'gridGap':17,'sectionGap':28,'playerHorizontalInset':22},
'radius':{'windowVisual':20,'thumbnail':14,'menu':16,'dialog':22},
'player':{'centralButton':88,'skipButton':62,'controlGap':26,'primaryIcon':39,'skipIcon':27,'autoHideSeconds':2.5,'preserveAspectRatio':True,'defaultCrop':False},
'window':{'minimum':[900,600],'testSizes':[[960,640],[1280,800],[1440,900]],'nativeTitlebar':True},
'behavior':{'cardSingleClick':'select','cardDoubleClick':'play','returnToBrowse':'continue session and show now-playing bar','historySaveIntervalSeconds':5},
'appearance':{'v02':'light','darkModeVersion':'v0.4','reduceTransparency':'opaque high-contrast fallback','reduceMotion':True},
'prototypeExceptions':['示例素材不是实际视频','示例属性需改为真实探测','目录入口需补全真实子目录','浏览时补底部当前播放栏','演示切换器与模拟按钮不进入原生产品']}
(OUT/'设计规范.json').write_text(json.dumps(tokens,ensure_ascii=False,indent=2))
(OUT/'开发任务.json').write_text(json.dumps({'schemaVersion':1,'planVersion':'1.1','date':'2026-10-13','tasks':tasks},ensure_ascii=False,indent=2))
(OUT/'调研来源.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2))
shutil.copy2(ROOT/'交付文件/Mac播放器_市场调研_PRD_技术方案_版本TODO.docx',OUT/'背景研究_2026年10月12日.docx')
shutil.copy2(UI/'assets/mountain-tram.jpg',ASSETS/'示例画面.jpg')
(ASSETS/'素材说明.txt').write_text('示例画面为 AI 生成的虚构场景，仅供 UI 开发和演示。不是用户的视频，不代表真实 H.265 播放结果。其他缩略图为设计稿内 CSS 示意图。实际产品使用用户自己的本地视频帧或合法授权的封面。\n')

def escaped(x):
    x=html.escape(x)
    return re.sub(r'\[(S\d+)\]',lambda m:f'<a href="{source_map[m[1]]["url"]}" target="_blank" rel="noreferrer">[{m[1]}]</a>',x)
parts=[];toc=[]
for kind,value in blocks:
    if kind=='h':
        anchor='ch'+value.split()[0];toc.append((anchor,value));parts.append(f'<h2 id="{anchor}">{escaped(value)}</h2>')
    elif kind=='s':parts.append('<h3>'+escaped(value)+'</h3>')
    elif kind=='p':parts.append('<p>'+escaped(value)+'</p>')
    elif kind=='b':parts.append('<p class="bullet">'+escaped(value)+'</p>')
    elif kind=='c':parts.append('<pre>'+escaped(value)+'</pre>')
    elif kind=='t':
        heads,rows=value;parts.append('<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+escaped(x)+'</th>' for x in heads)+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+escaped(x)+'</td>' for x in row)+'</tr>' for row in rows)+'</tbody></table></div>')
    elif kind=='image':
        name,caption=value;data=base64.b64encode((ASSETS/name).read_bytes()).decode();parts.append(f'<figure><img src="data:image/png;base64,{data}" alt="{escaped(caption)}"><figcaption>{escaped(caption)}</figcaption></figure>')
toc.extend([('tasks','9 逐版本可勾选任务'),('sources','10 调研和接口来源')])
parts.append('<h2 id="tasks">9 逐版本可勾选任务</h2><p>勾选只记录本机任务标记，不代表自动完成验收。导出保留依赖与验收内容。任务顺序以依赖为准。</p>')
for ver in version_names.values():
    parts.append('<h3>'+ver+'</h3>')
    for task in [x for x in tasks if x['version']==ver]:
        parts.append(f'<article class="task" data-id="{task["id"]}"><label><input type="checkbox" aria-label="标记任务 {task["id"]}"><strong>{task["id"]} {escaped(task["title"])}</strong></label><p>{escaped(task["acceptance"])}</p><small>{task["priority"]} · {task["size"]} · 依赖 {", ".join(task["dependsOn"]) or "无"} · 未开始</small></article>')
parts.append('<h2 id="sources">10 调研和接口来源</h2>')
for source in sources:parts.append(f'<p class="source"><strong>{source["id"]}</strong> {source["type"]}<br><a href="{source["url"]}" target="_blank" rel="noreferrer">{escaped(source["title"])}</a></p>')
css='''*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f0f4f8;color:#141d2d;font:15px/1.85 -apple-system,BlinkMacSystemFont,"PingFang SC",sans-serif}nav{position:fixed;width:260px;inset:0 auto 0 0;background:#fff;padding:30px 24px;border-right:1px solid #e5eaf0;overflow:auto}nav h1{font-size:25px;margin:0 0 12px}nav p{font-size:12px;color:#647084}nav a{display:block;text-decoration:none;padding:8px 0;color:#34445a;font-size:13px}nav a:hover{text-decoration:underline}main{margin:30px 30px 60px 290px;padding:45px 48px;background:white;border-radius:20px;max-width:1120px}header h1{font-size:32px;line-height:1.4;margin:0 0 16px}header p{color:#647084}h2{font-size:24px;scroll-margin-top:25px;margin:48px 0 18px}h3{font-size:18px;margin:30px 0 12px}p{margin:13px 0}a{color:#35617e;text-underline-offset:3px}.bullet{padding-left:17px;position:relative}.bullet:before{content:"•";position:absolute;left:0}.table-wrap{overflow:auto;margin:20px 0}table{border-collapse:collapse;width:100%;font-size:13px;min-width:550px}th,td{padding:11px 12px;vertical-align:top;text-align:left;border-bottom:1px solid #e5eaf0}th{background:#f0f4f8;white-space:nowrap}td:first-child{min-width:90px}pre{font:12px/1.8 Menlo,"PingFang SC",monospace;white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f7fa;border-radius:12px;padding:18px}figure{margin:20px 0}figure img{width:100%;display:block;border-radius:10px}figcaption{font-size:12px;color:#647084;margin-top:10px}.task{border-bottom:1px solid #e5eaf0;padding:16px 0}.task label{display:flex;gap:10px;align-items:baseline}.task input{accent-color:#141d2d;width:16px;height:16px}.task p,.task small{margin-left:26px}.task p{font-size:13px;margin-top:8px;margin-bottom:6px}.task small{font-size:11px;color:#647084}.task.done strong{text-decoration:line-through;color:#647084}.source{font-size:12px;border-bottom:1px solid #e5eaf0;padding:10px 0}button,.action{border:0;background:#141d2d;color:white;padding:10px 13px;border-radius:100px;font:12px -apple-system,sans-serif;margin:5px 4px 5px 0;cursor:pointer;text-decoration:none;display:inline-block}button:hover,.action:hover{background:#34445a}#count{font-size:12px;color:#647084;margin-top:14px}@media(max-width:950px){nav{position:static;width:auto;padding:22px}nav a{display:inline-block;margin-right:15px}main{margin:20px;padding:25px}header h1{font-size:27px}}@media print{nav{display:none}main{margin:0;padding:0;max-width:none}body{background:white;font-size:11pt}table{min-width:0;font-size:9pt}h2{break-before:page}h3{break-after:avoid}.task,tr,figure{break-inside:avoid}a{color:inherit;text-decoration:none}}'''
taskdata=json.dumps(tasks,ensure_ascii=False).replace('</','<\\/')
script='''const tasks=__DATA__,key='mori-dev-final-1.1';let marked={};try{marked=JSON.parse(localStorage.getItem(key)||'{}')}catch{}function update(){document.getElementById('count').textContent=Object.values(marked).filter(Boolean).length+' / '+tasks.length+' 项已标记'}document.querySelectorAll('.task').forEach(el=>{const box=el.querySelector('input');box.checked=!!marked[el.dataset.id];el.classList.toggle('done',box.checked);box.addEventListener('change',()=>{marked[el.dataset.id]=box.checked;el.classList.toggle('done',box.checked);try{localStorage.setItem(key,JSON.stringify(marked))}catch{}update()})});update();document.getElementById('export').onclick=()=>{const data={schemaVersion:1,planVersion:'1.1',tasks:tasks.map(t=>({...t,readerMarked:!!marked[t.id]}))};const u=URL.createObjectURL(new Blob([JSON.stringify(data,null,2)],{type:'application/json'}));const a=document.createElement('a');a.href=u;a.download='Mori开发任务.json';a.click();setTimeout(()=>URL.revokeObjectURL(u),1000)};document.getElementById('print').onclick=()=>print();'''.replace('__DATA__',taskdata)
page='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Mori macOS 播放器开发定稿</title><style>'+css+'</style></head><body><nav><h1>mori<br>开发定稿</h1><p>版本 1.1 · 2026 年 10 月 13 日<br>macOS 14+ · arm64 · 本地优先</p>'+''.join(f'<a href="#{a}">{escaped(v)}</a>' for a,v in toc)+'<a class="action" href="UI设计稿.html" target="_blank">查看 UI 设计稿</a><button id="export">导出任务</button><button id="print">打印</button><div id="count"></div></nav><main><header><h1>Mori macOS 播放器开发定稿</h1><p>产品需求 · 已确认 UI · 原生技术方案 · 版本任务 · 发布验收</p><p>先完成可靠的本地播放，再把目录、收藏与续播连接成一个完整的 Mac 使用流程。本文为开发执行基线；UI 为设计参考，软件功能尚待实现。</p></header>'+''.join(parts)+'<footer><p>本地方案与设计资源可直接保存到项目文档目录。以实现、真实测试和发布包验收作为完成依据。</p></footer></main><script>'+script+'</script></body></html>'
(OUT/'Mori播放器开发定稿.html').write_text(page)

# Editable core handbook. The full task records remain in the web edition and JSON.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for st in doc.styles:
    for border in list(st.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
for name,size in [('Normal',11),('Title',25),('Subtitle',12),('Heading 1',17),('Heading 2',13),('List Bullet',11)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB');st.paragraph_format.space_after=Pt(6);st.paragraph_format.line_spacing=1.16
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True;doc.styles[name].paragraph_format.space_before=Pt(12)
doc.styles['Normal'].paragraph_format.widow_control=True
doc.add_paragraph('Mori macOS 播放器开发定稿','Title')
doc.add_paragraph('产品需求  界面规范  原生架构  开发路线','Subtitle')
doc.add_paragraph('版本 1.1    2026 年 10 月 13 日    macOS 14 及以上    arm64')
doc.add_paragraph('以已确认的 UI 风格开发 Mac 本地播放器，首个公开版完成打开、播放、目录浏览、收藏和续播的完整流程。采用 Swift 原生 UI 与成熟播放内核，无需自建业务服务器。')
doc.add_paragraph(f'阅读入口：本手册为核心执行说明；网页方案提供完整的 {len(tasks)} 项可勾选任务与引用；开发任务.json 用于导入项目管理工具；UI设计稿.html 和设计规范.json 用于原生界面实现。')
def addlink(par,label,url):
    el=OxmlElement('w:hyperlink');el.set(qn('r:id'),par.part.relate_to(url,RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'315B77');rp.append(color);r.append(rp);tx=OxmlElement('w:t');tx.text=label;r.append(tx);el.append(r);par._p.append(el)
def rich(par,value):
    pos=0
    for m in re.finditer(r'\[(S\d+)\]',value):par.add_run(value[pos:m.start()]);addlink(par,m[0],source_map[m[1]]['url']);pos=m.end()
    par.add_run(value[pos:])
for kind,value in blocks:
    if kind=='h':doc.add_paragraph(value,'Heading 1')
    elif kind=='s':doc.add_paragraph(value,'Heading 2')
    elif kind in ['p','b']:rich(doc.add_paragraph(style='List Bullet' if kind=='b' else 'Normal'),value)
    elif kind=='c':
        for line in value.splitlines():
            par=doc.add_paragraph();par.paragraph_format.line_spacing=1.05;par.paragraph_format.space_after=Pt(2);r=par.add_run(line);r.font.name='Menlo';r.font.size=Pt(9);r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    elif kind=='image':
        name,caption=value;par=doc.add_paragraph();par.paragraph_format.keep_with_next=True;par.add_run().add_picture(str(ASSETS/name),width=Inches(6.8));par=doc.add_paragraph(caption);par.paragraph_format.space_after=Pt(10)
        for run in par.runs:run.font.size=Pt(9)
    elif kind=='t':
        heads,rows=value;tab=doc.add_table(rows=1,cols=len(heads));tab.autofit=False
        weights=[1.1]+[2.8]*(len(heads)-1)
        if len(heads)==4:weights=[.8,2,2.4,1]
        widths=[w/sum(weights)*7 for w in weights]
        for col,w in zip(tab.columns,widths):col.width=Inches(w)
        for cell,v in zip(tab.rows[0].cells,heads):cell.text=v
        repeat=OxmlElement('w:tblHeader');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,v in zip(tab.add_row().cells,row):rich(cell.paragraphs[0],v)
        for ri,row in enumerate(tab.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]);pr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
                for side,num in [('top',80),('bottom',80),('left',90),('right',90)]:el=OxmlElement('w:'+side);el.set(qn('w:w'),str(num));el.set(qn('w:type'),'dxa');margins.append(el)
                pr.append(margins)
                if ri==0:shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'EDF1F5');pr.append(shade)
                borders=OxmlElement('w:tcBorders');bt=OxmlElement('w:bottom');bt.set(qn('w:val'),'single');bt.set(qn('w:sz'),'4');bt.set(qn('w:color'),'DFE5EC');borders.append(bt);pr.append(borders)
                for par in cell.paragraphs:
                    par.paragraph_format.space_after=Pt(2);par.paragraph_format.line_spacing=1.12
                    for run in par.runs:run.font.size=Pt(9.5);run.font.bold=ri==0
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
doc.add_paragraph('9 任务索引与来源','Heading 1')
doc.add_paragraph(f'完整的 {len(tasks)} 项任务保存在开发任务.json 和网页方案第 9 章。按固定 ID 查找需求、依赖和验收；工程任务均未开始，视觉确认不替代实现状态。')
for ver in version_names.values():
    items=[x for x in tasks if x['version']==ver];doc.add_paragraph(f'{ver}  {len(items)} 项任务','Heading 2');doc.add_paragraph('；'.join(x['id']+' '+x['title'] for x in items)+'。')
for sid in ['S19','S20','S21','S25','S27','S32','S33','S34','S35','S36','S37','S38']:
    source=source_map[sid];par=doc.add_paragraph();par.add_run(sid+'  ');addlink(par,source['title'],source['url'])
doc.add_paragraph('完整市场背景和全部引用见背景研究及调研来源.json。所有能力以之后固定版本的实际测试和发布支持清单为准。')
footer=sec.footer.paragraphs[0];footer.alignment=2;r=footer.add_run('Mori 开发定稿  1.1    ');r.font.size=Pt(8);field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Mori macOS 播放器开发定稿';doc.core_properties.author='';doc.core_properties.subject='Mac 本地播放器产品界面架构与开发任务'
doc.save(OUT/'Mori播放器开发手册.docx')
(OUT/'开始开发.txt').write_text(f'''Mori macOS 播放器开发交付包
版本 1.1  日期 2026-10-13

1. 先打开 Mori播放器开发定稿.html，按第 1 至 8 章理解范围与实现方案。
2. 打开 UI设计稿.html，切换页面和菜单查看已确认视觉。设计为示例交互，不是真实播放器。
3. 设计规范.json 是原生 token 起点。原生 UI 补齐底部播放栏、真实子目录和权限状态。
4. 开发任务.json 含 {len(tasks)} 项任务、依赖和逐项验收；网页第 9 章可勾选并导出标记。
5. Mori播放器开发手册.docx 是可编辑的核心方案，完整任务在网页和 JSON。
6. 背景研究只保留市场判断和研究过程。UI、首版范围及交互冲突时，以 1.1 定稿为准。

第一步：建 Swift 6 原生工程，macOS 14+ arm64，固定 libmpv/FFmpeg 构建，播放一个真实 H.265 文件。验证声音、SDR 颜色、定位、停止释放和 UI 叠层。通过 POC 再接目录和 SQLite。

技术栈：SwiftUI + AppKit + libmpv + SQLite/GRDB + Sparkle 2。
服务器：无自建业务后端。更新与可选元数据/网络流访问外部服务，核心播放离线可用。
分发：静态官网 / GitHub Releases，Developer ID 签名公证，自带依赖 DMG。
已确认：视觉方向和产品路线。待完成：全部原生工程、播放测试与安装发布。
''')
ids={x['id'] for x in tasks};assert len(ids)==len(tasks)
def walk(tid,path):
    assert tid not in path,('cycle',path,tid)
    for dep in next(x for x in tasks if x['id']==tid)['dependsOn']:
        assert dep in ids,(tid,dep);walk(dep,path+[tid])
for tid in ids:walk(tid,[])
(ROOT/'final-build-check.json').write_text(json.dumps({'tasks':len(tasks),'sources':len(sources),'blocks':len(blocks),'docx':str(OUT/'Mori播放器开发手册.docx'),'html':str(OUT/'Mori播放器开发定稿.html')},ensure_ascii=False,indent=2))
print(json.dumps({'tasks':len(tasks),'sources':len(sources),'blocks':len(blocks),'out':str(OUT)},ensure_ascii=False))
