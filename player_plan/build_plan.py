from pathlib import Path
import json, re, html

ROOT = Path(__file__).resolve().parent
OUT = ROOT / '交付文件'
OUT.mkdir(exist_ok=True)
blocks = []
tasks = []
def h(text): blocks.append(('h1',text))
def s(text): blocks.append(('h2',text))
def p(text): blocks.append(('p',text))
def b(text): blocks.append(('b',text))
def t(headers, rows, widths=None): blocks.append(('table',(headers, rows, widths)))
def code(text): blocks.append(('code',text.strip()))
def todo(id, title, acceptance, depends='', size='M', priority='P0'):
    item=dict(id=id,title=title,acceptance=acceptance,depends=depends,size=size,priority=priority,status='未开始')
    tasks.append(item); blocks.append(('todo',item))

sources = [
('S01','IINA 1.5.0 发布说明','https://github.com/iina/iina/releases/tag/v1.5.0','官方版本记录'),
('S02','IINA 插件系统','https://iina.io/plugin/','官方功能'),
('S03','IINA 项目','https://github.com/iina/iina','官方项目'),
('S04','Infuse 官方功能','https://firecore.com/infuse','官方功能'),
('S05','Movist Pro 官方功能','https://movistprime.com/en/','官方功能'),
('S06','VidHub 正版应用页及评论','https://apps.apple.com/us/app/vidhub-video-library-player/id1659622164','开发者声明及用户反馈'),
('S07','SenPlayer 官网','https://www.senmuse.com/','官方功能'),
('S08','VideoLAN 项目介绍','https://images.videolan.org/projects/','官方功能'),
('S09','MacLC 官网','https://www.hazenstudio.com/maclc/','项目方声明'),
('S10','MacLC 开源仓库','https://github.com/OZBM/MacLC','官方项目'),
('S11','Owl 开源仓库','https://github.com/limboy/owl','官方项目'),
('S12','HorangPlayer 官网','https://www.horangplayer.com/','项目方声明'),
('S13','Cinema Player 开源仓库','https://github.com/hessennasser/cinema-player','官方项目'),
('S14','Infuse 用户关于文件夹结构的讨论','https://community.firecore.com/t/hello-brand-new-user-pro-purchase-attempting-to-use-with-plex/59291','个别用户反馈'),
('S15','IINA HDR 问题汇总','https://github.com/iina/iina/issues/6164','问题记录 不代表当前全部仍可复现'),
('S16','IINA 动态发现同目录文件请求','https://github.com/iina/iina/issues/6258','个别需求线索'),
('S17','Mac 玩家使用场景讨论','https://www.reddit.com/r/macbookpro/comments/1ikrt6q/macbook_pro_users_which_media_player_do_you/','个别用户反馈'),
('S18','mpv 官方能力介绍','https://mpv.io/','开发组件'),
('S19','mpv 渲染 API','https://github.com/mpv-player/mpv/blob/master/include/mpv/render.h','接口源码'),
('S20','mpv 控制 API','https://github.com/mpv-player/mpv/blob/master/include/mpv/client.h','接口源码'),
('S21','FFmpeg 格式及编解码器支持','https://ffmpeg.org/general.html#Supported-File-Formats_002c-Codecs-or-Features','官方技术文档'),
('S22','mpv 构建与许可说明','https://github.com/mpv-player/mpv','官方项目'),
('S23','SwiftUI 与 AppKit 集成','https://developer.apple.com/documentation/swiftui/appkit-integration','苹果开发文档'),
('S24','AVPlayerView','https://developer.apple.com/documentation/avkit/avplayerview','苹果开发文档'),
('S25','Apple 网站直接分发应用','https://developer.apple.com/documentation/technologyoverviews/distribution','苹果分发规则'),
('S26','Apple Developer Program 会员','https://developer.apple.com/programs/whats-included/','苹果会员说明'),
('S27','Sparkle 更新文档','https://sparkle-project.org/documentation/','官方技术文档'),
('S28','TMDB API FAQ','https://developer.themoviedb.org/docs/faq','服务方规则'),
('S29','KSPlayer 项目及许可','https://github.com/kingslay/KSPlayer','候选组件'),
('S30','Tauri 架构','https://tauri.app/concept/architecture/','候选组件'),
('S31','Mole 发布页','https://github.com/tw93/mole/releases','分发参考'),
('S32','mpv 完整使用文档','https://mpv.io/manual/stable/','官方技术文档'),
]

h('1 项目判断与实施决策')
p('结论：值得做一个限定范围的 Mac 开源播放器，先验证用户是否愿意持续使用，再决定长期投入。项目的成立条件是把漂亮的媒体库、直接打开文件的效率和可定制的播放交互结合好。单靠 Apple TV 外观、免费开源或广泛格式支持，吸引力不够持久。')
p('确定范围：只做 macOS；视觉参考 Mac 版 Apple TV App；处理普通本地视频和后续加入的标准网络流；通过官网和 GitHub Releases 下载；不以商业收入和 App Store 上架为目标。H.265 是常规编码支持的一部分。')
p('建议先投入 2 周验证播放与渲染，再用约 6 至 8 周做出可公开试用的 v0.2。熟悉 Swift 和原生 Mac 开发的一人团队，完整 v1.0 的规划量约 22 至 28 个全职工作周，包含约 30% 返工和兼容性缓冲。时间是规划估算，不能用来承诺发布日期。')
s('1.1 必须成立的三个用户价值')
b('随手打开：Finder 双击、拖入文件、拖入文件夹都能立即播放，不要求先建立媒体库，也不弹出多步引导。')
b('认真收藏：用户愿意添加目录时，同一批文件可以按原目录或电影与剧集浏览；离线硬盘保留目录和进度，重新连接后恢复。')
b('轻度可玩：观影界面保持克制；字幕、快捷键、控件位置和播放方案可调，设置可预览、撤销、导出和分享。')
s('1.2 为什么新应用而不是继续做插件')
p('IINA 插件已经能控制播放、访问 mpv、提供侧栏和视频覆盖层。因此，只做双字幕、快捷键、截图、书签或滤镜，优先给现有项目贡献插件会更划算。[S02]')
p('独立应用的理由是统一设计媒体库、文件夹浏览、播放窗口和返回路径。这些改变涉及应用主导航、持久化模型、窗口生命周期和播放器状态，超出了一个小插件的合理边界。建议新建薄应用外壳，复用播放内核；先不 fork 并长期维护整个 IINA 或 VLC。')
s('1.3 继续与收缩的条件')
t(['阶段','继续投入的证据','未达到时的动作'],[
['v0.0','本地 H.264 与 H.265 画面和声音正常；窗口切换稳定；原生界面覆盖层不卡顿','修复渲染和播放，不扩展媒体库'],
['v0.2','至少 12 人完成任务测试；至少 8 人完成核心任务；至少 6 人在第 2 周自愿继续使用','先改打开和返回路径，暂停在线信息与主题'],
['v0.4','至少 10 位持续用户，其中至少 5 人主动使用并认可核心差异；维护量可控','聚焦文件夹播放器，或把成果做成 IINA 扩展'],
['v1.0','兼容样本与回归测试通过；严重缺陷清零；网站下载到首次播放路径成立','保留 beta 状态，不为版本号强行发布']],[0.9,3.1,2.2])
p('上述人数是小规模验证门槛，不是市场统计或增长预测。GitHub Star、安装包下载量只能衡量关注度，不能替代真实播放和持续使用。')

h('2 市场调研与具体差异')
p('调研日期为 2026 年 10 月 12 日。检索覆盖主流 Mac 播放器、原生开源项目、官方发布记录、开发接口和公开用户讨论。产品功能比较依据官方资料；“用户为什么迁移”和“本项目能否更顺手”属于需要后续任务测试验证的假设。新项目的 README 和宣传页能证明已有相似方向，不能证明其稳定性或用户规模。')
s('2.1 主流产品已经解决的事情')
t(['产品','已经具备的强项','本项目可争取的体验'],[
['IINA 1.5','原生 Mac 播放、mpv、插件；最新发布已改进布局和加入插件商店 [S01][S03]','围绕个人收藏构建统一主界面，同时保留直接打开文件的短路径；不能再以“原生、插件、漂亮”宣称独有'],
['VLC','长期积累的格式与网络输入覆盖，开源生态 [S08]','更贴合 Mac 的浏览与播放闭环；格式广泛本身不是优势'],
['Infuse','精美媒体库、元数据、网络源和 Apple 多设备体验 [S04]','聚焦单台 Mac 的本地与目录工作流，提供可导出的交互配置；不能把海报墙或无广告当作新发明'],
['Movist Pro','字幕控制、同目录播放、缩略图、网络源及可选解码方式 [S05]','让收藏与播放处在同一导航体系，让常用进阶设置在播放中可理解、可回退'],
['VidHub','媒体库、NAS 与服务器连接、广泛格式；已有双字幕和界面迭代 [S06]','减少单机使用的设置成本，做可调的信息密度和一致的桌面快捷操作'],
['SenPlayer','影片媒体库与多类媒体源，覆盖 Mac [S07]','仅服务 Mac 的键鼠与窗口习惯；不扩展为网盘连接器集合']],[0.85,2.55,2.8])
s('2.2 不能忽视的新开源竞争者')
t(['项目','公开定位与能力','对我们意味着什么'],[
['MacLC','Mac 原生界面、媒体库、继续观看；基于 VLC 的开源项目，处于 public beta [S09][S10]','“免费开源的漂亮媒体库播放器”已有直接竞争；需要实测它的体验，不能假设市场空白'],
['Owl','SwiftUI 与 libmpv，文件夹媒体库、进度、双字幕和时间轴预览；网站式 DMG 分发 [S11]','“文件夹优先”也已经被实现；差异必须落到文件夹与收藏视图切换、返回路径和可组合交互'],
['HorangPlayer','原生 SwiftUI、mpv、滤镜、书签和循环 [S12]','原生 UI 加播放工具组合不是壁垒'],
['Cinema Player','SwiftUI 与 AVKit，本地优先、视频链接、可搜索媒体库 [S13]','轻量本地媒体库有相近开源方向；我们需要同时证明格式覆盖和使用流畅度']],[0.9,2.7,2.6])
p('建议在 v0.0 用同一组文件和任务实测 IINA、Infuse、MacLC、Owl。只比较真实完成路径，不采信各家“最快”“全格式”“零开销”等宣传口径。缺少可安装构建的项目只保留为设计参考。')
s('2.3 用户反馈给出的线索')
b('VidHub 应用页可见用户要求更合理的海报密度、Mac 原生感、更多媒体信息和批量已观看操作。这支持“信息密度与管理效率值得验证”，不代表所有用户都不满意。[S06]')
b('Infuse 论坛中有用户希望在服务器内容之外继续保留原有文件夹结构。应允许用户在目录和整理后的视图间来回切换。[S14]')
b('IINA 的 HDR 问题汇总说明颜色和显示器切换需要专门验证；历史问题不能直接作为当前版本仍有故障的证据。[S15]')
b('公开讨论中既有只想用文件夹管理的用户，也有重视原生外观的用户。产品应允许轻量打开，而不是强制整理所有内容。[S17]')
s('2.4 本项目的具体差异化承诺')
t(['拟做的体验','具体行为','与市场的关系','验证方式'],[
['播放器和收藏库共存','双击文件直播放；加入收藏后使用同一播放器；返回原目录和原滚动位置','各家有部分能力，本项目验证完整路径的一致性','5 个任务对比完成时间、误操作和偏好'],
['目录与内容两种视图','在原文件夹与电影剧集视图间切换，保持选中项和进度','Owl 擅长目录，Infuse 擅长内容整理；我们连接两个视角','跨视图找回同一文件，最多 2 次操作'],
['可调的信息密度','舒适与紧凑布局、海报大小、原文件名显示独立切换','可调布局已有竞争，需要做得可发现且一致','小屏与大屏对比，用户不靠猜即可调整'],
['可分享的播放方案','字幕、倍速、控件、快捷键分开成方案；一键预览和撤销','IINA 已有插件与配置；我们的差异是假设中的低门槛组合','新用户 60 秒内应用并恢复默认'],
['离线媒体不消失','外接硬盘拔出时保留封面和进度，重新接入恢复，不误删记录','成熟库产品有类似能力；作为基本可靠性要求','拔插、更名、目录移动后结果一致']],[1.05,1.85,1.8,1.4])
p('这些组合差异可以吸引偏好 Mac 原生体验和本地收藏的用户，但易被复制。长期优势将来自细节一致、缺陷响应、兼容性样本和社区贡献，而不是一种语言或一张主界面截图。')
s('2.5 产品表达和获取首批用户')
p('推荐对外表达：一款给 Mac 本地视频使用的开源播放器，可以随手打开，也可以认真收藏。功能说明以可演示的行为为依据，避免“所有格式”和“画质超过所有播放器”的承诺。')
p('官网首屏展示一段约 30 秒的真实操作录像：Finder 打开 MKV → 返回收藏库 → 切到原文件夹 → 恢复进度 → 调整双字幕 → 还原默认。首页的主要动作是下载，旁边是源码、系统要求和版本记录。Mole 可作为网站与开源发布路径的参考，不照搬其终端安装方式。[S31]')
p('先邀请 12 至 20 位已有 IINA 或 Infuse 使用习惯的 Mac 用户自愿试用，记录失败路径和迁移原因。待核心体验成立后，再发布到 V2EX、相关社区和 GitHub。方案不包含代用户发帖、创建账号或发布软件的动作。')

h('3 产品需求文档')
s('3.1 用户与目标场景')
t(['用户','当前行为','我们优先解决的任务'],[
['随手看视频的人','从 Finder、下载目录或外接硬盘打开单个视频','立即看、暂停、拖动、切字幕、下次继续'],
['本地收藏用户','文件夹里有电影、剧集与个人视频','保留原目录，同时有清晰的封面、进度和检索'],
['喜欢调节体验的人','经常调整字幕、倍速、快捷键或窗口','设置能被发现、预览、撤销并复用']],[1.3,2.1,2.8])
p('优先顺序为播放可靠性、打开与返回路径、目录和收藏浏览、轻度定制、网络输入。音乐库、剪辑工具、云盘聚合和多设备同步不进入首个稳定版。')
s('3.2 范围和版本边界')
p('首个公开体验版 v0.2 必须具备本地文件与目录播放、单窗口导航、Apple TV 风格首页、原目录视图、继续观看、搜索、字幕音轨和签名公证后的直接下载。电影自动匹配、进阶定制和标准网络流后续加入。')
p('v1.0 目标为个人 Mac 上的可靠本地播放器和小型媒体库，兼容经验证的主要视频、声音与字幕格式，并提供标准 HTTP/HLS/DASH 网络输入。不做跨平台、账号体系、内容商店、服务器端转码、自动下载站解析、蓝光菜单和大型脚本插件平台。')
p('应用只读引用用户媒体，索引、封面和进度写入自己的数据目录。“从媒体库移除”只移除引用。首版不提供原视频删除、移动或重命名能力。')
s('3.3 信息架构')
code('''主窗口
  首页：继续观看 / 最近加入 / 固定的目录或收藏
  视频：全部内容 / 电影 / 剧集 / 个人视频
  文件夹：用户添加的目录，保留原层级
  收藏：手工列表和收藏项
  搜索：文件名 / 手工标题 / 已缓存的作品标题
  底部当前播放栏：返回播放 / 暂停 / 下一项

播放界面
  主画面 + 临时控制浮层
  按需打开：播放队列 / 音轨字幕 / 画面与声音设置
  返回：保留浏览视图、滚动位置和选中项

设置
  常规 / 外观 / 播放 / 字幕 / 快捷键 / 数据 / 更新
''')
s('3.4 关键用户流程')
b('首次打开：窗口显示“打开视频”和“添加文件夹”；不要求账号或网络。打开文件后直接进入播放，媒体库稍后再建。')
b('Finder 双击：应用冷启动或已运行都能接收文件。默认沿用当前播放窗口；连续收到多文件时只改变队列一次，不反复抢焦点。')
b('添加目录：选择后立即展示已发现的文件名，封面和信息渐进补齐；扫描可取消；添加目录不复制原视频。')
b('浏览与播放：浏览单击选中，双击播放；卡片上明确的播放按钮也可直接播放。详情入口与播放入口分开，避免每次都强制经过详情。')
b('返回浏览：按返回按钮或 Esc 退出播放器的沉浸页面，播放器继续或暂停由用户设置决定；默认继续播放并显示当前播放栏。回到进入前的目录与滚动位置。')
b('继续观看：点击进度卡从保存位置开始；卡片菜单可“从头播放”和“标为已看”。短视频与长视频使用不同的完成判断，但用户可以覆盖结果。')
b('硬盘离线：卡片变为不可播放状态而非消失；点击说明来源离线，并提供重新定位目录。重新接入后自动核对，不重复添加内容。')
s('3.5 UI 还原与交互规格')
p('视觉方向为 Mac 版 Apple TV 的侧栏、内容分组、海报与横幅、克制的层级和播放浮层。界面有自己的名称、图标和用户内容。开始设计时固定 macOS 与 Apple TV App 版本，采集 960×640、1280×800、1440×900 三组参考窗口；这才是后续像素对照的基准。')
t(['项目','初始设计规格','验收'],[
['布局','侧栏初值 220 pt，可收起；内容边距 24 至 40 pt；最小窗口 900×600','连续缩放无重叠、裁切或横向溢出；紧凑窗口可收侧栏'],
['卡片','电影海报 2:3；普通视频与继续观看 16:9；缺图使用本地帧和清晰标题','不同来源不会被拉伸；无网络也有可读卡片'],
['排版','系统字体，标题正文元信息三级；时间数字等宽','中文长文件名、混合语言、200% 辅助文字不破坏操作'],
['动效','初值 160 至 240 ms；只对导航、展开、选中变化使用','减少动态效果启用时简化；动效不阻挡点击'],
['播放浮层','鼠标静止约 2.5 秒隐藏；聚焦控件、菜单展开或拖动中保持可见','退出拖动立即同步状态；不会误隐藏正在操作的控件'],
['参考还原','按选定 Apple TV 参考截图对比布局、字重、材质和层级','关键锚点误差目标不超过 4 pt；真实内容差异不按像素硬比'],
['系统外观','明暗模式；系统材质；macOS 26 增强效果可选','macOS 14/15 使用稳定替代材质，减少透明度时可读']],[1.1,2.7,2.4])
p('上表数值是待原型测试调整的设计初值，不是对 Apple TV 内部设计参数的测量结论。“1:1”应体现在参考基准、组件和交互的逐项核对，不能只靠堆叠玻璃效果。')
s('3.6 功能需求与验收')
t(['编号','需求','版本','核心验收'],[
['R01','单文件与多文件打开','0.1','拖入含空格和中文路径的文件可播放；多文件形成稳定有序队列'],
['R02','基本播放和音画同步','0.1','暂停、定位、倍速、音量和全屏状态正确；连续播放不漂移'],
['R03','字幕和音轨','0.1','内嵌与外部字幕可切换；音轨切换不中断 UI；字幕偏移可复位'],
['R04','进度和继续观看','0.1','正常退出后恢复误差目标小于 2 秒；异常退出最多损失最近一次保存间隔'],
['R05','文件夹与增量扫描','0.2','1 万条目录记录可检索；扫描可取消；离线和权限失败不误清空媒体库'],
['R06','统一浏览和播放','0.2','切回目录保留滚动与选择；浏览时继续播放不重新创建内核'],
['R07','首页和收藏','0.2','离线首页可用；固定目录、收藏、最近加入与继续观看行为一致'],
['R08','搜索和信息密度','0.2','中文文件名和标题能找到；密度变化不重置浏览上下文'],
['R09','作品整理和多版本','0.3','同作品不同文件可明确选择；手工改名不影响原文件'],
['R10','可选在线信息','0.3','关闭后停止网络匹配；失败不阻止播放；手工修正不被覆盖'],
['R11','播放方案与主题参数','0.4','方案可预览、撤销、导入导出；文件只包含声明式配置'],
['R12','双字幕和片段书签','0.4','双字幕位置独立；书签跳转正确；设置范围能被理解'],
['R13','网络视频与流媒体','0.5','HTTP/HLS/DASH 分别完成兼容矩阵；直播没有虚构总时长'],
['R14','网站下载和更新','0.2 起','干净 Mac 不安装 Homebrew 即可运行；升级不丢库与配置'],
['R15','数据与隐私','0.2 起','本地默认，无埋点；支持导出清除；隐藏来源可不留历史'],
['R16','辅助使用和本地化','0.2 起','主要操作支持键盘和 VoiceOver；简中、英文文案完整']],[0.6,1.5,0.7,3.4])
s('3.7 设置作用范围')
p('播放设置有四个层级：应用默认、用户方案、单文件记忆、当前会话临时调整。界面显示正在修改的层级；临时调整不会悄悄修改默认。字幕轨道与偏移默认按文件保存；目录的排序与密度按视图保存；音量为全局设置且切文件不突然重置。')
p('v0.4 的可玩性限定为颜色主题参数、布局参数、快捷键、字幕与播放方案、书签和循环。导入方案前展示变更摘要；应用后可撤销到导入前快照。首版不运行外部脚本。')
s('3.8 非功能需求')
p('所有性能数值都是拟定验收目标。基准设备为 M1 8 GB 与一台较新 Apple Silicon Mac；覆盖 macOS 14、15、26 可获得的真机或测试环境。缺少的设备组合不能标记为已验证。')
t(['指标','v0.2 目标','测量条件'],[
['启动到可操作','冷启动 P95 ≤ 2 秒','安装后已完成系统首次验证；本地 SSD；20 次重复'],
['本地首帧','H.264/H.265 1080p P95 ≤ 1.5 秒；4K P95 ≤ 2.5 秒','普通时长、完整索引、本地 SSD；记录首帧和首声'],
['定位','普通本地视频 P95 ≤ 800 ms','关键帧与文件索引影响单独报告；20 个定位点'],
['同步','30 分钟播放无持续漂移；目标误差 ≤ 80 ms','专用同步样片及录屏/测量流程，不能只听主观感觉'],
['大库交互','1 万条元数据下检索 P95 ≤ 200 ms；滚动 95% 帧 ≤ 16.7 ms','图片缓存热启动；同时扫描时分别记录'],
['内存','浏览 1 万条元数据目标 ≤ 350 MB；1080p SDR 播放总量目标 ≤ 600 MB','缓存有上限；4K/HDR 另报；连续切换 100 个文件无单调增长'],
['能耗','相同内容与设置，相对同版 IINA 不持续高于约 20%','先校准解码和渲染；取系统能耗指标，不凭 CPU 单项判定'],
['可靠性','所有 P0 样本通过；100 次切换无崩溃；2 小时播放通过','对固定测试包记录版本、系统、硬件及实际路径']],[1.0,2.5,2.7])
s('3.9 小规模用户验证')
p('用同一组普通文件在现用播放器和原型完成五个任务：打开 MKV；找到并继续昨天的视频；在目录与海报之间找回同一文件；调整字幕并恢复；播放中返回浏览再继续。记录完成率、时间、错误次数和理由，不要求用户先夸外观。')
p('成功假设：12 人中至少 8 人不经指导完成任务；至少 6 人愿意把原型用于接下来两周；至少 5 人能主动说出一个稳定的交互优势。若用户只是觉得截图好看，核心差异尚未成立。用户研究数据由用户主动提交，不默认收集使用事件。')

h('4 详细技术方案')
s('4.1 技术选型与取舍')
t(['部分','采用','理由与边界'],[
['语言与应用','Swift 6，Xcode 稳定版；macOS 14+','只做 Mac，便于窗口、系统控件、辅助功能与原生界面集成；新系统效果按可用性降级'],
['UI','SwiftUI 为主，AppKit 管理播放宿主和窗口','SwiftUI 与 AppKit 官方可互操作；高频渲染不经过 SwiftUI 状态树 [S23]'],
['播放器','libmpv 固定版本与构建配置','复用解码、字幕和播放能力；只封装自己的接口，不把 mpv 属性到处传给 UI [S18]'],
['渲染','v0.0 验证 libmpv OpenGL 嵌入；保留替换接口','公开 render API 当前列出 OpenGL 与软件路径，不能承诺现成 Metal 嵌入或原生 HDR [S19]'],
['数据','SQLite，使用成熟 Swift 封装，版本迁移','显式事务与索引，适合增量目录和 1 万条记录；SwiftData 不是必须'],
['媒体探测与预览','随应用分发 ffprobe/ffmpeg，后台辅助进程','与播放内核对齐版本；超时、取消和并发上限；前台播放优先'],
['更新','Sparkle 2，HTTPS appcast，EdDSA 更新签名','适合网站分发；更新签名与 Apple 代码签名分别验证 [S27]'],
['Rust','首版不作为必需依赖','Swift 已能承担目录、索引与后台任务；只有明确测量出瓶颈或共享需求才引入']],[1.05,2.1,3.05])
p('候选内核对照：AVPlayer/AVKit 的系统集成最直接，但仅靠系统格式覆盖不足以满足本项目；libVLC 可以作为 libmpv 验证失败后的备选；KSPlayer 方向适合 Apple 平台，但仓库当前有不同许可与部分付费能力，应按可获取的代码和功能逐项评估，而不是照抄宣传能力。[S24][S29]')
p('Tauri 可用于 Rust 与网页界面组合，但 Mac 单平台目标没有充分理由增加 WebView 与原生视频混合层。本方案不使用浏览器 video 标签作为广泛格式播放底座。[S30]')
s('4.2 模块结构与依赖方向')
code('''AppShell / Navigation / WindowCoordinator
              ↓
UI ViewModels 与 AppState             LibraryStore
              ↓                           ↓
PlaybackCoordinator                  LibraryRepository
     ↓             ↓                   ↓        ↓
PlaybackEngine  RenderSurface       SQLite     ScanService
     ↓             ↓                              ↓
MPVAdapter      AppKit HostView      Probe / Preview Helper
     ↓                                            ↓
libmpv + 固定依赖                         ffprobe / ffmpeg

横向服务：SettingsStore / HistoryStore / UpdateService / Diagnostics
后续可选：MetadataProvider / StreamSource / ProfileStore
''')
p('领域对象保持纯 Swift 值类型；UI 不直接接触 mpv_handle、文件描述符和数据库连接。PlaybackCoordinator 决定队列与意图，MPVAdapter 只负责内核命令和事件，WindowCoordinator 只负责窗口与视频视图的归属。目录扫描和图片生成不能阻塞播放控制。')
s('4.3 播放接口')
code('''protocol PlaybackEngine {
    func open(_ request: OpenRequest) async throws
    func play() async throws
    func pause() async throws
    func seek(to seconds: Double, mode: SeekMode) async throws
    func selectAudio(_ id: TrackID) async throws
    func selectSubtitle(_ id: TrackID?) async throws
    func setRate(_ rate: Double) async throws
    func stop() async
    var events: AsyncStream<PlaybackEvent> { get }
    var capabilities: PlaybackCapabilities { get }
}

OpenRequest：requestID、source、resumeAt、selectedTracks、profile
PlaybackEvent：loading、ready、position、buffering、tracksChanged、ended、failed
PlaybackCapabilities：seek、live、dualSubtitle、HDRMode、hardwareDecode
''')
p('接口示意以职责为准。实现时 RenderSurface 单独管理平台资源，不要求它跨 actor 传递。每次打开生成 requestID；晚到的旧请求事件必须丢弃。定位请求合并，只保留用户最近一次意图。任何失败都返回可解释的错误分类和恢复动作。')
s('4.4 状态机与生命周期')
code('''idle → opening → ready → playing ⇄ paused
                  │          │
                  └→ failed  ├→ buffering → playing / failed
                             ├→ seeking → playing / paused
                             └→ ended → next / idle

新 open 请求：取消旧准备流程，保留新 requestID
用户 stop：保存当前进度，停止输出，再释放媒体资源
切换全屏或返回媒体库：只改变呈现，不重建播放内核
退出应用：停止事件投递 → 释放渲染上下文 → 销毁内核
''')
p('playing 与 paused 是用户意图；buffering、seeking 是暂态，不能把网络恢复后的播放状态猜成自动播放。保存队列索引与当前文件 ID，结束事件需识别自然结束、用户停止和加载失败，防止误跳下一集。')
s('4.5 线程和渲染的约束')
p('内核控制与事件消费使用独立串行执行域；C 回调只唤醒事件泵，不直接改 SwiftUI，不在回调中同步等待 UI。Swift actor 不保证固定 OS 线程，不能仅靠 actor 就假设 OpenGL context 始终在正确线程。')
p('渲染使用专用执行域，遵守 context current 与 mpv 渲染线程要求；控制线程与渲染线程之间避免相互等待。mpv 的文档明确提示相关 API 混用可能产生死锁。[S19][S20]')
p('UI 位置更新合并到约 10 Hz，画面渲染按内核帧更新处理；拖动时本地时间显示实时响应，释放后发最终精确定位。不要每帧转成 SwiftUI Image，不通过主线程拷贝全分辨率 RGBA。')
p('NSViewRepresentable 包装一个生命周期稳定的 AppKit 视频宿主。主播放页和当前播放栏之间迁移同一个宿主视图或改变它的布局，确保任一时刻只有一个父视图和一个有效渲染上下文。SwiftUI 重算不能隐式创建第二个 mpv 实例。最小化、锁屏和屏幕变化分别测试。')
s('4.6 HDR 颜色与原画质')
p('原文件保持只读，正常播放只有解码与显示，不执行重新编码。“不改原文件”不等于“显示颜色必然正确”：色彩空间、全幅与限幅、显示器 ICC、HDR 与 SDR 映射都需要验证。')
p('v0.0 先确保 SDR 的 H.264/H.265 在内置与外接显示器上的范围、灰阶和色彩正确。v0.1 的确定承诺是 SDR；HDR 输入只有经验证的 SDR 映射路径才列为支持，未通过就如实显示限制。真正的 EDR/HDR 输出不作为早期 UI 版本的前置卖点。')
p('同时比较当前 IINA 与 QuickTime 对兼容样片的呈现，使用标准测试图和硬件信息解释差异。主观截图不能证明 HDR 正确，系统截图本身也可能经过映射。[S15]')
p('若原生 HDR 是用户验证后的高优先级需求，单独开 2 至 4 周验证 AVFoundation 路径或另一渲染内核。保持 PlaybackEngine 接口，但不在首版同时维护多个自动切换内核。切换内核可能改变字幕、轨道、截图和进度行为，需要完整兼容测试。')
s('4.7 数据模型')
t(['表或实体','关键字段','约束与用途'],[
['Source','id、kind、rootBookmark、volumeID、state、lastScanGeneration','表示一个目录或后续网络源；离线不是删除'],
['MediaItem','id、type、displayTitle、sortTitle、manualOverrides','表示用户看到的内容；允许 unknown、movie、episode、personal'],
['MediaFile','id、itemID、sourceID、relativePath、resourceID、fingerprint、size、mtime、availability','同一作品可关联多文件；unique(sourceID, relativePath)'],
['MediaTrack','fileID、kind、nativeTrackID、language、codec、title','音轨字幕随探测结果更新；轨道 ID 仅在对应文件内解释'],
['PlaybackState','fileID、position、duration、watchedOverride、trackChoices、updatedAt','进度按文件保存，不把不同时长版本直接套用同一时间'],
['Collection 与 Member','collectionID、itemID、sortOrder','手工收藏和列表；删除集合不删除内容'],
['MetadataMatch','itemID、provider、providerID、confidence、provenance、lastFetch','在线信息有来源；手工覆盖优先，低信心不自动改名'],
['Bookmark','fileID、position、label、thumbnailKey','片段书签，与媒体文件身份绑定'],
['Profile','id、schemaVersion、type、payload','声明式播放与外观方案；导入前验证'],
['CacheEntry','key、kind、bytes、lastAccess、sourceFingerprint','数据库只存缓存索引；实际图像在 Caches']],[1.0,2.8,2.4])
p('数据关系：Source 1 对多 MediaFile；MediaItem 1 对多 MediaFile；MediaFile 1 对多 Track、Bookmark，并有一条当前 PlaybackState；Collection 与 MediaItem 多对多。结构允许后续添加网络源，但不预建复杂服务器、账号或多租户体系。')
p('文件身份不能只用绝对路径，也不能只相信 inode。先用来源 ID、volume/resource 标识和相对路径；更名或移动时结合大小、修改时间与后台抽样指纹匹配。检测到标识重用或内容替换时生成新身份。默认不对整部大视频做全量哈希。')
p('数据库版本迁移有递增 schemaVersion，使用事务；迁移前创建本地备份。升级失败保留旧库，进入只读恢复或提示恢复，不能清空重建掩盖失败。导出包含逻辑 ID、来源引用、进度、手工信息与方案，不含媒体文件本身。')
s('4.8 扫描与离线恢复')
p('扫描分三步：枚举可见文件；读取快速基础属性；按用户可见性和最近播放优先探测、生成缩略图。首屏先出现文件名与占位图。遍历不跟随无限符号链接，去重嵌套来源，限制并发。')
p('本地目录用文件事件触发去抖后的增量核对；事件丢失、目录变更或挂载变化时做局部或完整对账。只有来源在线、权限有效、扫描完整成功时，才能把本轮未发现项标为 missing。取消、超时、权限错误和磁盘断开都不能触发批量清理。')
p('在应用重启时解析保存的目录 bookmark，检查失效并引导重新定位。非 App Store 分发也仍要尊重文件访问权限与系统隐私限制，不要求用户授予整个磁盘访问权。将书签访问与来源生命周期绑定，避免泄漏授权句柄。')
s('4.9 缩略图与在线信息')
p('预览进程最多同时两个任务，播放开始时缩减到一个或暂停。每任务有超时和取消；用参数数组启动命令，不拼接 shell 字符串。通过固定可执行路径运行，清理临时文件，回收被取消的子进程。')
p('卡片图优先本地 sidecar 图片，其次已缓存在线海报，再次本地视频抽帧。时间轴预览按需生成，不为每个文件一次生成整条时间轴。缓存 key 包含文件指纹、抽帧时间、尺寸和处理版本；LRU 上限初值 1 GB，可调并可清理。')
p('在线信息默认关闭，由用户对特定目录启用。匹配时只发送用于查询的标题与年份；不上传视频或全路径。先读取本地 NFO 和图片，再做可选 TMDB 匹配；手工选择与修正有最高优先级。服务端故障时使用本地卡片继续播放。')
p('TMDB 接入需要按其规则申请凭据并提供署名；“非商业开源”不等于没有服务使用约束。桌面二进制内的公共 API key 不能当作可保密的服务器凭据；配置可更换并控制请求速率，不把个人私钥放进仓库。[S28]')
s('4.10 网络输入的边界')
p('v0.5 先实现标准 HTTP/HTTPS 文件 URL、HLS、DASH。URL 先分类，凭据保存在 Keychain；日志移除 query token、认证头和用户目录。网络模式与本地模式复用播放界面，但 capability 不同。')
p('VOD 根据真实总时长提供定位；直播没有确定总时长就禁用普通进度条，显示直播和缓存状态。没有 HTTP Range 的资源不能假装可随机定位，也不默认把整部大视频偷偷缓存到磁盘。')
p('HLS/DASH 能打开不代表自动拥有与网页流媒体引擎相同的自适应码率能力。v0.5 明确验证默认变体选择、手动切换、定位和断线行为；没有验证动态 ABR 就不提供误导性的“自动随网速切换”选项。[S21][S32]')
p('SMB 首版通过 Finder 挂载后的目录使用，产品表现为普通文件来源。WebDAV、Jellyfin、Emby、Plex 的浏览和账号接入后置评估；它们不是简单填入播放 URL 就完成的功能。RTSP 等输入也只在有样本和真实需求后追加。')
s('4.11 可玩性和扩展设计')
p('所有内置动作登记到 CommandRegistry：命令 ID、标题、可用条件、快捷键和处理入口共用。菜单、键盘和后续命令面板调用同一动作，避免不同入口的行为漂移。')
p('Profile JSON 有 schemaVersion、明确字段白名单和数值范围。主题只允许颜色、间距密度、圆角与可见控件配置；不得引用外部可执行文件、脚本或任意本地路径。导入提供字段级预览，保存旧快照以撤销。')
p('运行脚本的插件平台至少等 v1.0 用户验证后再决定。若未来要提供文件或网络能力，必须有清楚的权限与隔离设计；JavaScript 放在同一进程运行不自动构成安全隔离。当前版本的可玩性无需承担整个插件生态的维护。')
s('4.12 网站分发与自动更新')
p('推荐路径：静态官网 → GitHub Releases 的签名公证 DMG → 拖入 Applications → 首次打开 → 直接播放。安装包自带播放依赖，普通用户不需要 Homebrew、终端和额外解码包。')
p('不上 App Store 仍建议使用 Developer ID 签名和 Apple 公证。公证属于自动检查，不是 App Store 审核。Apple Developer Program 当前通常为每年 99 美元或当地货币；这是友好分发的维护成本，与收费商业模式无关。[S25][S26]')
code('''固定源码与依赖清单
 → 测试与 Release 构建
 → 收集 dylib 并改为应用内相对加载路径
 → 检查架构、rpath、依赖与许可
 → 从内到外签名辅助程序、库和 App
 → Hardened Runtime 与实际需要的最小 entitlement
 → 打包 DMG / 更新 ZIP
 → notarytool 公证 → staple 票据 → 验证 Gatekeeper
 → 生成 SHA256 与 Sparkle EdDSA 签名
 → 准备 GitHub Release、源码归档和 appcast
 → 在干净 Mac 下载验证后，人工批准正式发布
''')
p('CI 只准备构建与发布候选包。签名私钥放在专用 Keychain 或受保护的 CI secret，不能进入日志。更新分 stable/beta 两个渠道；应用在播放期间只提示，不强制退出安装。关闭更新检查后不主动联网检查。')
p('官网下载按钮链接具体架构和最低系统要求，包含变更记录、已知限制、SHA256、源码与反馈入口。先发布 arm64；如果拿不到 Intel 真机测试，不提供写着“已支持 Intel”的通用包。用户在 Mac 平台内部的 Intel 需求可以作为后续版本事项。')
s('4.13 开源与依赖管理')
p('建议应用采用 GPL-3.0-or-later，并对实际随包依赖逐项核对兼容性。mpv 默认 GPLv2-or-later，也提供关闭 GPL 的 LGPL 构建方式；实际许可由构建及依赖决定，不能给所有 mpv 包统一贴 MIT 标签。[S22]')
p('每个发布版本保留自身源码、依赖版本与源码快照、补丁、构建脚本、配置参数、第三方许可和 notices。网站链接对应发布版本的源码与构建说明。初版不 fork 播放内核；需要补丁时尽量先提交上游，并将补丁数量保持可控。')
p('支持清单来自随包二进制实际探测结果，加上测试矩阵；不是从 FFmpeg 官网复制一份列表就声称全部支持。更新内核前重跑媒体样本、颜色、字幕、窗口和能耗回归。')
s('4.14 兼容性矩阵和样本管理')
t(['类型','首批 P0 样本','后续或明确限制'],[
['封装','MP4、MOV、MKV、WebM、AVI、TS 的常用组合','FLV、WMV、RMVB 等按样本追加；ISO/BDMV 与菜单不作为 v1 承诺'],
['视频','H.264；H.265 Main/Main10；VP9；AV1 的常用样本','AV1 能否硬解依硬件；H.265 4:2:2、12-bit 和特殊 profile 单列'],
['声音','AAC、MP3、PCM、FLAC、Opus；AC3/EAC3、DTS 的解码样本','解码成 PCM 不等于 Atmos 或 TrueHD 位流直通'],
['字幕','UTF-8 SRT、ASS、VTT；MKV 内嵌字幕；常见中文编码','PGS 等图片字幕需要样本；强制字幕、复杂 ASS 样式逐项验证'],
['时间与显示','23.976/24/25/30/60 fps；可变帧率；旋转 MOV；长文件名','4K60 压力、HDR10/HLG、外接显示器分别报告'],
['网络','HTTP Range、HLS TS/fMP4、DASH 的 VOD 与直播','无 Range、断线、凭据过期和无时长必须有正确 UI'],
['操作','切换音轨字幕、快速 seek、拔盘、睡眠恢复、全屏与多屏','mini window 与 PiP 分开验收，不把置顶窗口叫系统 PiP']],[1.0,2.7,2.5])
p('v0.0 建立约 30 个基础样本，v0.1 扩到约 60 个，v1.0 扩到至少 120 个有明确元数据的样本。包含公开许可视频和自制测试片，不在仓库上传商业电影。manifest 记录 codec、封装、时长、授权来源、期望轨道、预期表现和校验值。')
p('测试分层：纯逻辑单测覆盖排序、身份、迁移、完成判断与方案校验；集成测试覆盖真实内核打开和事件；UI 测试覆盖五条核心用户流程；手工硬件测试覆盖 HDR、多屏、蓝牙声音与能耗。视频画面和声音的真实呈现不能只靠 mock。')

h('5 版本路线与逐项 TODO')
p('任务规模 S 为约半天至 1 天，M 为约 1 至 2 天，L 为约 3 至 5 天；仅表示相对规模。每个版本先完成 P0，再做 P1。所有任务当前均为未开始，文档不表示软件已经实现。依赖字段给出阻塞关系；同一版本中无显式依赖的 UI 与数据工作可顺序或交错推进。')
t(['版本','建议工作量','用户可感知结果','发布条件'],[
['v0.0 验证','约 2 周','播放、渲染与设计基准得到证据','内部原型，未通过就不扩展'],
['v0.1 播放器','约 3 周','可靠打开文件与基本播放','小范围 alpha'],
['v0.2 公开体验','约 3 至 4 周','Apple TV 风格首页与目录收藏闭环','第一份公证 DMG 和任务测试'],
['v0.3 内容整理','约 2 至 3 周','作品与剧集、可选元数据','离线仍可用'],
['v0.4 个性化','约 2 至 3 周','可预览撤销的播放方案与进阶工具','用户确实使用差异功能'],
['v0.5 网络输入','约 2 至 3 周','HTTP/HLS/DASH 明确支持','流媒体矩阵与失败恢复通过'],
['v0.9 候选版','约 2 周','数据迁移与兼容性稳定','冻结功能，严重缺陷清零'],
['v1.0 稳定版','约 1 周','可信下载与长期使用基础','发布包、更新、源码和文档一致']],[1.2,1.0,2.4,1.6])
p('版本合计约 17 至 21 周直接实施；加上约 30% 渲染、兼容和设计返工，按约 22 至 28 个全职工作周规划。若每周只投入 12 至 15 小时，大致按 14 至 22 个月考虑；熟悉度和硬件测试条件会显著影响时间。v0.2 之前不做大量服务器连接和脚本插件。')
s('5.1 v0.0 播放与体验验证')
todo('A01','固定产品边界和五个对比任务','范围记录中只有 Mac、普通视频、网站分发；五个任务与用户邀请清单齐备',size='S')
todo('A02','记录 Apple TV App 视觉基准','三种窗口尺寸的首页、详情、播放、菜单和空态参考齐备；记录系统与 App 版本','A01','M')
todo('A03','实测四个核心竞品','同一文件包对比 IINA、Infuse、MacLC、Owl；记录版本、路径和失败，不填想象的评分','A01','M')
todo('A04','建立可分发媒体测试包','30 个有许可来源的样本及 manifest；包含 H.265 Main10 与同步片','A01','M')
todo('A05','建立 Swift 工程与内核依赖锁定','干净构建可运行；记录 libmpv/FFmpeg 来源和构建配置；有许可证目录',size='M')
todo('A06','实现最小播放与事件桥','H.264 与 H.265 有画面和声音；异步打开、停止和退出无挂死','A04,A05','L')
todo('A07','验证原生视频宿主与 UI 叠层','连续 100 次缩放、覆盖控件、全屏和窗口迁移没有重复内核或渲染崩溃','A06','L')
todo('A08','验证 SDR 色彩和 HDR 限制','内置与外接显示样片结果有记录；明确 SDR/HDR 映射与真实 HDR 的支持状态','A06,A07','M')
todo('A09','测量首帧能耗与停止资源释放','基线数据可复现；停止后 CPU/GPU 回落；连续切文件没有内存持续增长','A06','M')
todo('A10','完成继续开发决策','播放证据通过；确定默认渲染路径及未通过能力的降级，不以截图代替验证','A02,A03,A07,A08,A09','S')
s('5.2 v0.1 可用本地播放器')
todo('B01','实现 PlaybackEngine 与状态机','旧请求事件被丢弃；暂停、缓冲、定位和结束意图分离','A10','M')
todo('B02','实现 Finder 打开与拖放','单个、多文件、目录拖入，中文空格路径全部可用；应用已启动时不丢事件','B01','M')
todo('B03','实现 Apple TV 风格基本播放浮层','播放暂停、时间、进度、音量、全屏可见；操作期间不自动隐藏','A02,B01','M')
todo('B04','实现精确定位与拖动合并','快速连续拖动以最终位置为准；暂停时 seek 不自动播放','B01,B03','M')
todo('B05','实现字幕与音轨菜单','内嵌、外部字幕、语言与偏移可用；菜单显示真实轨道','B01','M')
todo('B06','实现播放队列与自然排序','S01E2 排在 S01E10 前；结束与手动停止不混淆；错误文件可恢复','B02','M')
todo('B07','实现数据层和播放进度','SQLite 事务与首个迁移；进度按文件保存；正常退出恢复误差目标小于 2 秒','B01','M')
todo('B08','实现窗口全屏快捷键和媒体键','原生全屏、Esc、空格、方向键、Now Playing 正常；焦点不互抢','B03','M')
todo('B09','实现可理解的失败界面','损坏、缺失、权限、解码失败分别提供动作；可回到浏览状态','B01','S')
todo('B10','提供设置和简中英文文案','默认行为清楚；所有核心操作可键盘完成；没有漏出的内部错误术语','B03,B05','M')
todo('B11','完成 60 样本与操作回归','P0 本地样本有声音和图像；30 分钟同步、100 次切换、睡眠恢复通过','B04,B05,B06,B07','L')
todo('B12','制作自带依赖的内部安装包','没有 Homebrew 的测试 Mac 能启动；dylib 路径全部可解析','A05,B11','M')
s('5.3 v0.2 第一份公开体验版')
todo('C01','实现 Source 与 MediaItem 文件身份','目录引用唯一；同一文件不会因简单更名丢进度；内容替换得到新身份','B07','L')
todo('C02','实现可取消渐进目录扫描','文件名先出现；取消、离线和权限失败不会清空现有库','C01','L')
todo('C03','实现增量更新与离线恢复','新增、删除、移动、拔插盘和事件丢失有核对；离线卡片保持可见','C02','M')
todo('C04','实现本地缩略图和受限缓存','抽帧超时可取消；最多两个后台任务；上限清理不会删除原视频','C02','M')
todo('C05','实现 Apple TV 风格首页和侧栏','继续观看、最近加入、固定目录与空态在三种尺寸下完成视觉验收','A02,C04','L')
todo('C06','实现文件夹与卡片密度视图','原层级保留；舒适紧凑切换不丢选择；超长标题能辨认','C02,C05','M')
todo('C07','实现搜索收藏和批量已看操作','1 万条库可检索；收藏可撤销；已看操作不改原文件','C01,C06','M')
todo('C08','完成播放与浏览的无重载切换','返回原滚动位置；当前播放栏与主播放器共用会话，无重复声音','B08,C05,C06','L')
todo('C09','完成隐私模式和数据管理','指定目录可不记历史；支持导出和清除应用记录；默认无在线匹配','C01,C07','M')
todo('C10','完成键盘辅助使用和大库性能','VoiceOver 主要路径可用；1 万条目标及内存指标有测试记录','C05,C06,C07','L')
todo('C11','完成 Developer ID 与公证发布流程','DMG 从浏览器下载后 Gatekeeper 验证通过；不需要用户执行终端命令','B12','M')
todo('C12','准备静态官网下载与版本说明','下载指向明确版本；架构、最低系统、源码、校验值与已知限制齐备','C11','M')
todo('C13','接入 Sparkle beta 更新并测试','旧 beta 到新 beta 升级成功；保留库和进度；播放中不强制更新','C11','M')
todo('C14','执行 12 至 20 人公开体验验证','五任务有记录，达到继续门槛；用户反馈对照竞品而非只问好不好看','C08,C09,C10,C12','L')
s('5.4 v0.3 内容组织与可选信息')
todo('D01','实现电影剧集个人视频分类','未知内容保留原文件名；分类可手改和撤销','C14','M')
todo('D02','实现同作品多版本与剧集排序','多个文件不被错误覆盖；播放前能选版本；季与集有自然顺序','D01','M')
todo('D03','读取本地封面和 NFO','侧边文件优先于网络；坏数据有上限和降级，不影响播放','D01','M')
todo('D04','实现目录级在线元数据开关','用户主动启用；关闭后停止匹配；超时限流失败都走本地内容','D03','M')
todo('D05','实现匹配修正和手工优先级','用户选中的标题海报不会被下一轮扫描覆盖；可恢复自动结果','D04','M')
todo('D06','完成电影详情和剧集详情','播放与详情入口分开；展示真实文件信息；返回保留位置','D02,D05','L')
todo('D07','实现目录与作品视图映射','同一文件两种视角可相互定位，进度与收藏保持一致','C06,D02,D06','M')
todo('D08','完成元数据缓存与署名','API 凭据可更换；有缓存清理；按服务规则展示来源署名','D04','S')
todo('D09','验证未匹配与大库场景','个人视频、乱码文件名、同名不同年作品与离线网络测试通过','D05,D07','M')
todo('D10','验证内容整理的真实价值','至少 5 位收藏用户完成找片、纠错和多版本选择；无持续误整理','D06,D09','M')
s('5.5 v0.4 可玩性与播放方案')
todo('E01','实现统一命令注册表','菜单、快捷键和 UI 按钮调用同一动作；不可用状态一致','C14','M')
todo('E02','实现快捷键编辑与冲突提示','可搜索动作；冲突可见；系统快捷键保持正确；能恢复默认','E01','M')
todo('E03','实现设置分层和播放方案','默认、方案、文件与会话四层可解释；临时改动不污染默认','D10,E01','L')
todo('E04','实现配置导入预览撤销与版本迁移','非法字段和越界值被拒绝；导入摘要清楚；撤销回到旧快照','E03','M')
todo('E05','实现有限外观参数与预设','密度、强调色和控件可见性可调；全部预设满足对比度与可读性','E04','M')
todo('E06','实现双字幕与独立位置设置','内嵌加外部组合正确；偏移独立；恢复默认不残留旧配置','B05,E03','M')
todo('E07','实现书签 A B 循环和逐帧控制','时间点保存正确；循环状态明显；退出循环操作易发现','E01,E03','M')
todo('E08','实现时间轴按需预览','停留后再抽帧；快速移动取消旧任务；不抢前台解码资源','C04','M')
todo('E09','实现可选置顶小窗','迁移同一播放宿主，返回主窗口不重载；不冒称系统 PiP','C08','L',priority='P1')
todo('E10','验证设置可发现性与使用率','新用户 60 秒应用并恢复方案；至少 5 位持续用户主动使用一个差异功能','E02,E04,E06,E07','M')
s('5.6 v0.5 标准网络输入')
todo('F01','实现 URL 输入与安全凭据保存','普通 URL 与认证分开；Keychain 存凭据；历史和日志不暴露 token','E10','M')
todo('F02','实现 HTTP 文件播放与 Range 行为','支持定位的资源正确 seek；无 Range 时如实降级且不偷偷全量下载','F01','M')
todo('F03','验证 HLS TS 与 fMP4 VOD','变体选择、字幕、首帧、seek 与结束行为有样本记录','F01','L')
todo('F04','验证 DASH VOD','依赖构建包含所需支持；音视频流同步，缺轨道有解释','F01','M')
todo('F05','实现直播呈现与能力判断','无总时长不显示假进度；仅有真实 DVR 窗口才允许回看','F03,F04','M')
todo('F06','实现超时重试与用户取消','有界重试；停止后结束连接；旧请求恢复不能抢占新视频','F02,F05','M')
todo('F07','实现源级缓存限制和状态','可见缓存与重连状态；磁盘和内存有上限；清理不影响原文件','F02,F06','M')
todo('F08','形成网络兼容清单与回归','弱网、断网、认证过期、无 Range、直播和 VOD 分开报告','F03,F04,F05,F06,F07','L')
todo('F09','测试 Finder 挂载 NAS 文件夹','睡眠与断开后源为 offline，重新挂载恢复；不扩展成账号连接器','C03,F08','M',priority='P1')
s('5.7 v0.9 发布候选版')
todo('G01','冻结功能并补齐兼容矩阵','至少 120 样本；通过和限制公开；无未解释的 P0 失败','F08','L')
todo('G02','核验色彩与显示器变化','SDR 范围与颜色通过；HDR 映射、真实 HDR 逐项标记；多屏无崩溃','A08,G01','L')
todo('G03','完成历史版本数据库迁移回滚测试','从每个公开版本升级；中断迁移不丢库；备份可恢复','D10,E04','M')
todo('G04','完成退出睡眠声音设备生命周期回归','2 小时播放；蓝牙与输出切换；停止及退出资源全部释放','G01,G02','M')
todo('G05','完成辅助功能和三尺寸 UI 回归','键盘、VoiceOver、明暗、减少透明与动态全部通过','E05,G01','M')
todo('G06','完成无 Homebrew 干净安装测试','两台可用目标 Mac 下载、公证、安装、升级、移除均通过','C13,G03','M')
todo('G07','完成依赖及日志清理','所有 dylib 与 helper 来源和许可齐备；日志默认不含路径凭据','G01,F01','M')
todo('G08','完成持续用户候选版验证','至少 10 位持续用户反馈，所有阻断问题处理；非核心请求留待以后','G04,G05,G06','M')
s('5.8 v1.0 稳定发布')
todo('H01','关闭阻断与严重缺陷','P0/P1 缺陷清零；其他缺陷公开且可绕过；发布清单有人复核','G08','M')
todo('H02','归档源码构建与兼容性证据','版本号、commit、依赖、源码、许可、构建配置和样本结果对应','G07,H01','S')
todo('H03','准备稳定更新渠道与发布包','stable 与 beta 分离；签名有效；校验和与 appcast 指向同一包','G06,H02','M')
todo('H04','完成官网支持与反馈入口','首次播放引导、FAQ、已知限制、问题模板和贡献文档齐备','C12,H02','M')
todo('H05','执行最终真实下载安装验收','浏览器下载 → 首次运行 → H.265 播放 → 重启续播 → 更新全链路通过','H03,H04','S')
todo('H06','决定后续版本优先级','根据持续用户需求选一个方向；不同时推进服务器连接、HDR 重写和插件平台','H05','S')
s('5.9 v1.1 以后待验证事项')
p('候选方向依次比较：真实 HDR/EDR 路径；Intel 构建与测试；系统 PiP；WebDAV 或一个媒体服务器连接；Quick Look 扩展；经过权限设计的插件系统。每次只选择一个方向，先收集样本与用户任务，再开独立版本。v1.0 不以实现所有候选作为完成条件。')

h('6 实施顺序与维护方式')
s('6.1 第一周工作安排')
t(['工作日','交付结果'],[
['第 1 天','冻结范围；建立五个用户任务；整理 Apple TV 界面参考'],
['第 2 天','收集授权样片；记录四个竞品的同任务行为'],
['第 3 天','Swift 工程、libmpv 链接与 C 接口包装；播放一个 H.265 文件'],
['第 4 天','AppKit 宿主与 SwiftUI 覆盖层；窗口缩放和全屏反复测试'],
['第 5 天','记录首帧、颜色、停止释放和明显风险；决定第二周验证内容']],[1.0,5.2])
s('6.2 个人开发者维护策略')
p('主线始终可构建，按版本建立小范围任务。每个发布版本有固定媒体样本与五条核心路径。功能请求按可重复任务排序，不按评论声音大小排序。播放问题优先拿到最小样本与系统信息，再判断属于应用、内核还是文件。')
p('按月检查播放依赖更新与已知问题，但不自动追所有上游最新版。依赖升级作为独立改动发布 beta。贡献入口优先欢迎字幕、翻译、兼容样本、键盘辅助和小范围 UI 修复，不把核心渲染交给未经验证的零散补丁。')
p('用户自愿导出诊断包时先展示包含内容，移除路径、文件名和网络 token。社区反馈可用匿名文件 ID 与媒体元信息定位问题。默认不发送遥测；持续使用指标通过主动参与测试取得。')
s('6.3 风险与降级')
t(['风险','触发条件','处理'],[
['渲染与 HDR 超出能力','内核嵌入路径颜色不正确，或窗口迁移不稳定','暂停扩展功能，优先通过 SDR；独立评估另一渲染路线'],
['差异过小','用户完成任务后仍认为现用播放器同样顺手','缩小为目录播放器，或把有价值功能贡献给上游'],
['媒体库拖慢播放','扫描抽帧导致音画卡顿或 UI 卡顿','暂停后台任务，降低并发，限制缓存，按可见内容加载'],
['资料服务不可靠','元数据 API 不可用、匹配错误或额度不足','默认本地可用；停用在线匹配仍有标题缩略图'],
['直接分发安装困难','签名、依赖或公证失败','修复发布流程；不把关闭系统保护当成正常安装步骤'],
['维护范围失控','同时要求云盘、服务器、插件、投屏和剪辑','保留路线图，按验证结果一次做一个方向']],[1.3,2.0,2.9])
s('6.4 最终建议')
p('开工，但把第一阶段限定为可比较、可放弃的验证。选择 Swift 原生 UI 与成熟播放内核；在 v0.2 交付完整的直接打开、目录浏览、收藏首页和返回路径。只有用户愿意持续使用，才继续投入电影整理、方案分享和网络输入。这个项目有做出鲜明体验的空间，空间来自完成路径和长期打磨，不能只靠 Apple TV 风格复刻。')

h('7 调研来源')
p('以下来源供核对产品功能、接口与分发规则。功能声明、公开问题和用户偏好是不同证据，引用时分别标注。检索日期为 2026 年 10 月 12 日。')
for sid,title,url,kind in sources:
    blocks.append(('source',(sid,title,url,kind)))

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

source_map = {x[0]: x for x in sources}
def link(par, label, url, color='245A81'):
    rel=par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    el=OxmlElement('w:hyperlink'); el.set(qn('r:id'),rel)
    r=OxmlElement('w:r'); rp=OxmlElement('w:rPr')
    c=OxmlElement('w:color'); c.set(qn('w:val'),color); rp.append(c)
    r.append(rp); tx=OxmlElement('w:t'); tx.text=label; r.append(tx);el.append(r);par._p.append(el)

def rich(par,text):
    pos=0
    for m in re.finditer(r'\[(S\d+)\]',text):
        par.add_run(text[pos:m.start()]); sid=m.group(1)
        link(par,m.group(0),source_map[sid][2]);pos=m.end()
    par.add_run(text[pos:])

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
for name,size in [('Normal',10.5),('Title',24),('Subtitle',12),('Heading 1',17),('Heading 2',12.5),('List Bullet',10.5)]:
    st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.line_spacing=1.22
for name in ['Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.space_before=Pt(13)
    doc.styles[name].paragraph_format.keep_with_next=True
doc.styles['Normal'].paragraph_format.widow_control=True
title=doc.add_paragraph('Mac 开源播放器产品与技术方案','Title')
doc.add_paragraph('市场调研  产品需求  技术设计  版本任务','Subtitle')
doc.add_paragraph('研究日期 2026 年 10 月 12 日    文档版本 1.0    平台 macOS')
doc.add_paragraph('建议做一个聚焦 Mac 本地视频体验的开源产品，先验证播放正确性与完整交互路径，再投入长期维护。本文给出研究依据、可实施需求、技术边界和分版本任务。')
doc.add_paragraph('阅读顺序：第 1 至 2 章用于判断是否开工；第 3 章为 PRD；第 4 章为技术方案；第 5 章为可执行 TODO；第 6 章为实施与维护；第 7 章为来源。')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Mac 播放器方案  ·  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='Mac 开源播放器产品与技术方案'
doc.core_properties.subject='Mac 播放器市场分析 PRD 技术架构与版本路线'
doc.core_properties.author=''

for kind,value in blocks:
    if kind=='h1':
        par=doc.add_paragraph(value,'Heading 1')
        par.paragraph_format.page_break_before=True
    elif kind=='h2': doc.add_paragraph(value,'Heading 2')
    elif kind in ('p','b'):
        par=doc.add_paragraph(style='List Bullet' if kind=='b' else 'Normal');rich(par,value)
    elif kind=='code':
        for line in value.splitlines():
            par=doc.add_paragraph();par.paragraph_format.space_after=Pt(1)
            par.paragraph_format.line_spacing=1.0
            r=par.add_run(line);r.font.name='Menlo';r.font.size=Pt(8.3)
            r._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Hiragino Sans GB')
    elif kind=='table':
        headers,rows,widths=value
        tab=doc.add_table(rows=1,cols=len(headers));tab.alignment=WD_TABLE_ALIGNMENT.CENTER
        tab.autofit=False
        if widths is None: widths=[1]*len(headers)
        total=sum(widths); widths=[w/total*6.65 for w in widths]
        for col,w in zip(tab.columns,widths):col.width=Inches(w)
        for i,v in enumerate(headers):tab.rows[0].cells[i].text=v
        repeat=OxmlElement('w:tblHeader');repeat.set(qn('w:val'),'true');tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=tab.add_row().cells
            for i,v in enumerate(row):
                cells[i].text='';rich(cells[i].paragraphs[0],v)
        for ri,row in enumerate(tab.rows):
            trp=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trp.append(cant)
            for ci,cell in enumerate(row.cells):
                cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for side in ['top','left','bottom','right']:
                    el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
                tcpr.append(borders);margins=OxmlElement('w:tcMar')
                for side,num in [('top',75),('bottom',75),('left',90),('right',90)]:
                    el=OxmlElement('w:'+side);el.set(qn('w:w'),str(num));el.set(qn('w:type'),'dxa');margins.append(el)
                tcpr.append(margins)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EDF1' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));tcpr.append(shade)
                for par in cell.paragraphs:
                    par.paragraph_format.space_after=Pt(1);par.paragraph_format.space_before=Pt(1)
                    par.paragraph_format.line_spacing=1.1
                    for rr in par.runs:rr.font.size=Pt(9.2);rr.font.bold=ri==0
        doc.add_paragraph().paragraph_format.space_after=Pt(2)
    elif kind=='todo':
        d=value
        par=doc.add_paragraph();par.paragraph_format.keep_with_next=True
        par.paragraph_format.space_after=Pt(2)
        r=par.add_run(f"{d['id']}  {d['title']}");r.bold=True
        par=doc.add_paragraph('验收：'+d['acceptance']);par.paragraph_format.keep_with_next=True;par.paragraph_format.space_after=Pt(2)
        par=doc.add_paragraph(f"优先级 {d['priority']}  ·  规模 {d['size']}  ·  依赖 {d['depends'] or '无'}  ·  未开始")
        par.paragraph_format.space_after=Pt(8)
        for r in par.runs:r.font.size=Pt(8.5);r.font.color.rgb=RGBColor.from_string('505860')
    elif kind=='source':
        sid,ttl,url,category=value
        par=doc.add_paragraph();par.paragraph_format.space_after=Pt(8)
        par.add_run(f'{sid}  {category}  ');link(par,ttl,url)
docpath=OUT/'Mac播放器_市场调研_PRD_技术方案_版本TODO.docx'
doc.save(docpath)

def hr(text):
    txt=html.escape(text)
    return re.sub(r'\[(S\d+)\]',lambda m:f'<a class="cite" href="{html.escape(source_map[m.group(1)][2])}" target="_blank" rel="noreferrer">[{m.group(1)}]</a>',txt)

parts=[];toc=[];opened=False;version=''
for kind,value in blocks:
    if kind=='h1':
        if opened:parts.append('</section>')
        anchor='chapter-'+value.split()[0];toc.append((anchor,value))
        parts.append(f'<section id="{anchor}"><h2>{hr(value)}</h2>');opened=True
    elif kind=='h2':
        version=value if value.startswith('5.') else version
        parts.append(f'<h3>{hr(value)}</h3>')
    elif kind=='p':parts.append(f'<p>{hr(value)}</p>')
    elif kind=='b':parts.append(f'<p class="bullet">{hr(value)}</p>')
    elif kind=='code':parts.append(f'<pre>{html.escape(value)}</pre>')
    elif kind=='table':
        heads,rows,_=value
        parts.append('<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+hr(x)+'</th>' for x in heads)+'</tr></thead><tbody>')
        parts.extend('<tr>'+''.join('<td>'+hr(x)+'</td>' for x in row)+'</tr>' for row in rows)
        parts.append('</tbody></table></div>')
    elif kind=='todo':
        d=value;d['version']=version
        parts.append(f'<article class="task" data-id="{d["id"]}"><label><input type="checkbox" aria-label="标记 {d["id"]}"><strong>{d["id"]} {hr(d["title"])}</strong></label><p>验收：{hr(d["acceptance"])}</p><small>{d["priority"]} · 规模 {d["size"]} · 依赖 {d["depends"] or "无"} · 计划任务</small></article>')
    elif kind=='source':
        sid,ttl,url,category=value
        parts.append(f'<p class="source"><b>{sid}</b> <span>{category}</span><br><a href="{html.escape(url)}" target="_blank" rel="noreferrer">{hr(ttl)}</a></p>')
parts.append('</section>')
css='''
:root{--ink:#17202a;--muted:#66717e;--line:#dce2e9;--accent:#245a81;--nav:#14202d}*{box-sizing:border-box}html{scroll-behavior:smooth}body{margin:0;background:#f1f4f7;color:var(--ink);font:16px/1.8 -apple-system,BlinkMacSystemFont,"PingFang SC",sans-serif}nav{position:fixed;inset:0 auto 0 0;width:260px;background:var(--nav);padding:36px 24px;color:#e6edf5;overflow:auto}nav .eyebrow{font-size:11px;letter-spacing:1.6px;color:#a2b3c6}nav h1{font-size:24px;line-height:1.4;margin:15px 0 8px}nav p{font-size:13px;color:#b3c0ce}nav a{display:block;color:#d8e2ec;text-decoration:none;font-size:14px;padding:9px 0;border-bottom:1px solid #293646}nav a:hover{color:white}nav .actions{margin-top:24px}button{border:1px solid #45617d;border-radius:7px;padding:9px 12px;background:#22364b;color:white;cursor:pointer;margin:6px 4px 6px 0;font:inherit;font-size:12px}main{margin:32px 32px 60px 292px;max-width:1040px;background:white;padding:48px 54px;border-radius:12px;box-shadow:0 3px 22px #10294008}header .eyebrow{color:var(--muted);font-size:13px}header h1{font-size:35px;line-height:1.35;margin:12px 0 18px}header p{color:#4b5865}.summary{font-size:19px;color:var(--ink);margin-bottom:36px}section{scroll-margin-top:28px}h2{font-size:27px;line-height:1.4;margin:55px 0 22px}h3{font-size:19px;margin:32px 0 12px;color:#182532}p{margin:12px 0}a{color:var(--accent);text-underline-offset:3px}.cite{font-size:.8em;white-space:nowrap}.bullet{padding-left:20px;position:relative}.bullet:before{content:"•";position:absolute;left:0;color:#557489}.table-wrap{overflow-x:auto;margin:20px 0}table{border-collapse:collapse;width:100%;min-width:600px;font-size:14px;line-height:1.65}th,td{border:1px solid var(--line);padding:12px 14px;text-align:left;vertical-align:top}th{background:#eaf0f5;font-weight:600}tbody tr:nth-child(even){background:#f8fafb}pre{white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f6f8;padding:20px;border-radius:8px;font:13px/1.65 ui-monospace,Menlo,"PingFang SC",monospace}.task{padding:18px 0;border-bottom:1px solid var(--line)}.task label{display:flex;gap:10px;align-items:baseline;cursor:pointer}.task input{accent-color:var(--accent);width:16px;height:16px}.task p,.task small{margin-left:26px}.task small{color:var(--muted);font-size:12px}.task.done strong{text-decoration:line-through;color:var(--muted)}.source{font-size:14px;padding:10px 0;border-bottom:1px solid #edf0f4}.source span{color:var(--muted);font-size:12px}.notice{font-size:12px!important;color:#aebdcc!important}#progress{font-size:13px;color:#d8e2ec;margin-top:16px}@media(max-width:1000px){nav{position:static;width:100%;padding:22px}nav a{display:inline-block;margin-right:14px}nav h1{font-size:22px}nav .actions{margin-top:8px}main{margin:20px;padding:28px}header h1{font-size:28px}}@media print{nav{display:none}main{margin:0;padding:0;box-shadow:none;max-width:none}body{background:white;font-size:10pt}h2{break-before:page}h3{break-after:avoid}table{min-width:0;font-size:8pt}tr,.task{break-inside:avoid}a{color:inherit;text-decoration:none}.task input{display:none}pre{font-size:8pt}header h1{font-size:24pt}}
'''
script='''
const tasks=TASK_DATA;
const key='mac-player-plan-2026-10-12-v1';let saved={};
try{saved=JSON.parse(localStorage.getItem(key)||'{}')}catch(e){}
const update=()=>{document.querySelector('#progress').textContent=`已标记 ${Object.values(saved).filter(Boolean).length} / ${tasks.length} 项`};
document.querySelectorAll('.task').forEach(el=>{const box=el.querySelector('input');box.checked=!!saved[el.dataset.id];el.classList.toggle('done',box.checked);box.onchange=()=>{saved[el.dataset.id]=box.checked;el.classList.toggle('done',box.checked);try{localStorage.setItem(key,JSON.stringify(saved))}catch(e){}update()}});update();
document.querySelector('#export').onclick=()=>{const data=tasks.map(t=>({...t,readerMarked:!!saved[t.id]}));const url=URL.createObjectURL(new Blob([JSON.stringify(data,null,2)],{type:'application/json'}));const a=document.createElement('a');a.href=url;a.download='Mac播放器_版本TODO.json';a.click();setTimeout(()=>URL.revokeObjectURL(url),1000)};
document.querySelector('#print').onclick=()=>window.print();
'''.replace('TASK_DATA',json.dumps(tasks,ensure_ascii=False).replace('</','<\\/'))
page='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Mac 开源播放器产品与技术方案</title><style>'+css+'</style></head><body><nav><div class="eyebrow">MAC PLAYER · RESEARCH & PLAN</div><h1>Mac 开源播放器<br>产品与技术方案</h1><p>2026 年 10 月 12 日<br>仅 macOS · Mac 版 Apple TV 风格</p>'+''.join(f'<a href="#{anchor}">{html.escape(label)}</a>' for anchor,label in toc)+'<div class="actions"><button id="export">导出版本 TODO</button><button id="print">打印或保存 PDF</button></div><div id="progress"></div><p class="notice">勾选仅记录本机阅读或任务进度，不代表功能已实现或验收。方案本身不是播放器安装包。</p></nav><main><header><div class="eyebrow">市场调研 / 产品需求 / 技术方案 / 版本 TODO</div><h1>Mac 开源播放器产品与技术方案</h1><p>研究日期 2026 年 10 月 12 日 · 文档版本 1.0</p><p class="summary">建议开工做一个有明确边界的验证版本。用成熟播放内核保证基础能力，把差异落到随手打开、目录与收藏共存，以及可预览撤销的个性化交互。</p></header>'+''.join(parts)+'<footer><p>文档与任务均为开发规划。实际能力以之后发布版本的测试与支持清单为准。</p></footer></main><script>'+script+'</script></body></html>'
htmlpath=OUT/'Mac播放器_完整方案.html';htmlpath.write_text(page,encoding='utf-8')
(OUT/'Mac播放器_版本TODO.json').write_text(json.dumps(tasks,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'调研来源.json').write_text(json.dumps([dict(id=a,title=b,url=c,type=d) for a,b,c,d in sources],ensure_ascii=False,indent=2),encoding='utf-8')
ids={x['id'] for x in tasks};assert len(ids)==len(tasks)
for task in tasks:
    for dep in filter(None,task['depends'].split(',')):assert dep in ids,(task['id'],dep)
def check_cycle(id,path):
    assert id not in path,('dependency cycle',path,id)
    item=next(x for x in tasks if x['id']==id)
    for dep in filter(None,item['depends'].split(',')):check_cycle(dep,path+[id])
for id in ids:check_cycle(id,[])
report={'tasks':len(tasks),'sources':len(sources),'blocks':len(blocks),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'text_characters':sum(len(str(x[1])) for x in blocks),'outputs':[str(docpath),str(htmlpath)]}
(ROOT/'build_check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))


